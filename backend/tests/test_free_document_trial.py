"""Real uploads, shared free budget, export rights, refunds and PostgreSQL race.

No translation provider or payment collection is called. Only the worker is
replaced; extraction, admission, persistence and authenticated downloads run.
"""
import asyncio
from datetime import timedelta
import io
from pathlib import Path
import unittest
from urllib.parse import urlsplit
import uuid

import racine  # noqa: F401
import test_beta_program as fixtures
from fastapi import Response
from docx import Document as WordDocument
from pptx import Presentation
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker
from sqlalchemy.pool import NullPool

from app.api import translate as translation
from app.core.database import DATABASE_URL
from app.models import User, Document, UsageRecord, KIND_FREE_OFFER, KIND_CREDITS
from app.services import usage
from app.services.documents import count_pages
from app.services.jobs import jobs


def docx(words):
    output = io.BytesIO()
    document = WordDocument()
    document.add_paragraph(" ".join(["bonjour"] * words))
    document.save(output)
    return output.getvalue()


def pptx():
    output = io.BytesIO()
    document = Presentation()
    slide = document.slides.add_slide(document.slide_layouts[0])
    slide.shapes.title.text = "Présentation de test"
    document.save(output)
    return output.getvalue()


class FreeDocumentTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.fx = fixtures.BetaTests()
        await self.fx.asyncSetUp()
        self.client = self.fx.client
        self.user = self.fx.other
        self.headers = self.fx.headers(self.user)

    async def asyncTearDown(self):
        await self.fx.asyncTearDown()

    async def upload(self, data=None, filename="document.pdf", **options):
        return await self.client.post("/api/translate", headers=self.headers,
            files={"file": (filename, data if data is not None else fixtures.pdf(), "application/octet-stream")},
            data={"target_lang": "en", "use_beta": "0", **options})

    async def summary(self):
        response = await self.client.get("/api/auth/me/usage", headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    async def document(self, response):
        self.assertEqual(response.status_code, 200, response.text)
        identifier = jobs.get(response.json()["job_id"])["document_id"]
        async with self.fx.sessions() as db:
            return await db.get(Document, identifier)

    async def test_all_five_formats_share_five_pages_as_preview_only(self):
        initial = await self.summary()
        self.assertEqual(initial["monthly_pages"], 5)
        self.assertEqual(initial["pages_remaining"], 5)
        for filename, data in [("a.pdf", fixtures.pdf()), ("b.pptx", pptx()),
                ("c.docx", docx(20)), ("d.xlsx", fixtures.xlsx(20)), ("e.txt", b"Bonjour le monde")]:
            response = await self.upload(data, filename)
            document = await self.document(response)
            # Pages gratuites : aperçu seulement, pas de téléchargement.
            self.assertFalse(response.json()["can_download"])
            self.assertFalse(document.paid)
            self.assertEqual(document.page_count, 1)
            result = Path(self.fx.temp.name) / filename
            result.write_bytes(data)
            jobs.set(response.json()["job_id"], state="done", result_path=str(result), result_filename=filename)
            download = await self.client.get("/api/translate/result/" + response.json()["job_id"], headers=self.headers)
            self.assertEqual(download.status_code, 402, filename)
        self.assertEqual((await self.summary())["pages_remaining"], 0)
        denied = await self.upload()
        self.assertEqual(denied.status_code, 402)
        self.assertEqual(denied.json()["pages_missing"], 1)
        self.assertEqual(len(self.fx.queued), 5)

    async def test_office_and_txt_costs_use_actual_text_and_reject_oversized_free_upload(self):
        for ext, data in [("docx", docx(501)), ("xlsx", fixtures.xlsx(501)),
                          ("txt", ("bonjour " * 501).encode()), ("txt", ("漢" * 501).encode())]:
            with self.subTest(ext=ext):
                self.assertEqual(count_pages(data, ext), 2)
        self.assertEqual(count_pages(docx(500), "docx"), 1)
        # A compressed file containing one enormous token cannot buy unlimited
        # model input by pretending it contains just one word.
        huge_token = await self.upload(("a" * 150001).encode(), "unbroken.txt")
        self.assertEqual(huge_token.status_code, 402, huge_token.text)
        self.assertGreater(huge_token.json()["pages_needed"], 5)
        response = await self.upload(fixtures.xlsx(2501), "large.xlsx", covered_only="1")
        self.assertEqual(response.status_code, 402, response.text)
        self.assertEqual(response.json()["pages_needed"], 6)
        self.assertEqual(response.json()["pages_missing"], 1)
        self.assertFalse(response.json()["partial_possible"])
        self.assertEqual((await self.summary())["pages_remaining"], 5)
        self.assertEqual(len(self.fx.queued), 0)

    async def test_free_and_purchased_budget_combine_and_refunds_are_idempotent(self):
        async with self.fx.sessions() as db:
            (await db.get(User, self.user.id)).page_credits = 2
            await db.commit()
        response = await self.upload(docx(2501), "large.docx")
        document = await self.document(response)
        # Document mixte : des pages achetées le rendent téléchargeable.
        self.assertTrue(response.json()["can_download"])
        self.assertTrue(document.paid)
        summary = await self.summary()
        self.assertEqual((summary["pages_remaining"], summary["page_credits"]), (0, 1))
        async with self.fx.sessions() as db:
            records = (await db.execute(select(UsageRecord).where(UsageRecord.document_id == document.id))).scalars().all()
            self.assertEqual(sorted((row.kind, row.pages) for row in records),
                             sorted([(KIND_FREE_OFFER, 5), (KIND_CREDITS, 1)]))
            for _ in range(2):
                await usage.refund_for_document(db, document.id)
                await db.commit()
        summary = await self.summary()
        self.assertEqual((summary["pages_remaining"], summary["page_credits"]), (5, 2))
        async with self.fx.sessions() as db:
            for _ in range(2):
                await usage.rebill_for_document(db, document.id)
                await db.commit()
        summary = await self.summary()
        self.assertEqual((summary["pages_remaining"], summary["page_credits"]), (0, 1))

    async def test_deletion_does_not_restore_quota_and_new_month_does(self):
        document = await self.document(await self.upload(fixtures.pdf(5)))
        async with self.fx.sessions() as db:
            await db.delete(await db.get(Document, document.id))
            await db.commit()
        self.assertEqual((await self.upload()).status_code, 402)
        async with self.fx.sessions() as db:
            record = (await db.execute(select(UsageRecord).where(UsageRecord.user_id == self.user.id))).scalar_one()
            self.assertIsNone(record.document_id)
            record.created_at = usage.month_start() - timedelta(seconds=1)
            await db.commit()
        self.assertEqual((await self.summary())["pages_remaining"], 5)
        self.assertEqual((await self.upload()).status_code, 200)

    async def test_partial_pdf_requires_explicit_choice_and_charges_only_selection(self):
        rejected = await self.upload(fixtures.pdf(7), pages="2-7")
        self.assertEqual(rejected.status_code, 402)
        self.assertEqual(rejected.json()["pages_missing"], 1)
        self.assertEqual(len(self.fx.queued), 0)
        accepted = await self.upload(fixtures.pdf(7), pages="2-7", covered_only="1")
        document = await self.document(accepted)
        self.assertEqual(document.page_count, 5)
        self.assertEqual(jobs.get(accepted.json()["job_id"])["pages"], {2, 3, 4, 5, 6})

    async def test_verified_account_required_and_bad_files_do_not_consume_quota(self):
        async with self.fx.sessions() as db:
            (await db.get(User, self.user.id)).email_verified = False
            await db.commit()
        self.assertEqual((await self.upload()).status_code, 403)
        async with self.fx.sessions() as db:
            (await db.get(User, self.user.id)).email_verified = True
            await db.commit()
        self.assertEqual((await self.upload(b"invalid zip", "bad.docx")).status_code, 422)
        self.assertEqual((await self.summary())["pages_remaining"], 5)
        self.assertEqual(len(self.fx.queued), 0)

    async def test_old_locked_previews_and_other_users_documents_stay_private(self):
        identifier = jobs.create()
        self.fx.queued.append(identifier)
        result = Path(self.fx.temp.name) / "old.pdf"
        result.write_bytes(fixtures.pdf())
        jobs.set(identifier, user_id=self.user.id, state="done", result_path=str(result), result_filename="old.pdf")
        response = await self.client.get("/api/translate/result/" + identifier, headers=self.headers)
        self.assertEqual(response.status_code, 402)
        response = await self.client.get("/api/translate/result/" + identifier, headers=self.fx.tester_headers)
        self.assertEqual(response.status_code, 404)
        response = await self.client.post("/api/translate", files={"file": ("file.pdf", fixtures.pdf())})
        self.assertEqual(response.status_code, 401)

    async def test_displayed_plan_totals_match_the_payment_quote_including_annual_rounding(self):
        for country in ["CM", "FR", "MA"]:
            response = await self.client.get("/api/pricing", params={"country": country})
            pricing = response.json()
            for plan in pricing["plans"]:
                if not plan["for_sale"]:
                    continue
                for annual in [False, True]:
                    quote = await self.client.post("/api/payments/plan/quote",
                        headers={**self.headers, "CF-IPCountry": country},
                        json={"plan": plan["key"], "annual": annual})
                    self.assertEqual(quote.status_code, 200, quote.text)
                    # Les cartes affichent le net : les frais n'apparaissent
                    # que dans la fenêtre de paiement.
                    net = plan["annual"] * 12 if annual else plan["monthly"]
                    self.assertEqual(net, quote.json()["subtotal"])
                    self.assertNotIn("monthly_total", plan)
            if country == "CM":
                starter, pro = pricing["plans"][1:3]
                self.assertEqual((starter["monthly"], pro["monthly"]), (1500, 4500))
                self.assertEqual((starter["annual"], pro["annual"]), (1200, 3500))


class PostgreSQLFreeQuotaRace(unittest.IsolatedAsyncioTestCase):
    @unittest.skipUnless(urlsplit(DATABASE_URL).hostname in ("localhost", "127.0.0.1")
                         and urlsplit(DATABASE_URL).path == "/precis", "isolated CI PostgreSQL only")
    async def test_parallel_three_page_uploads_cannot_spend_six_of_five_free_pages(self):
        engine = create_async_engine(DATABASE_URL, poolclass=NullPool)
        sessions = async_sessionmaker(engine, expire_on_commit=False)
        try:
            async with engine.connect():
                pass
        except Exception:
            await engine.dispose()
            self.skipTest("CI PostgreSQL not available")
        fx = fixtures.BetaTests()
        await fx.asyncSetUp()
        fx.sessions = sessions
        async with sessions() as db:
            user = User(email=f"free-race-{uuid.uuid4().hex}@example.com", email_verified=True, plan="free",
                        created_at=racine.compte_ancien())
            db.add(user)
            await db.commit()
        async def upload():
            async with sessions() as db:
                account = await db.get(User, user.id)
                return await translation.translate_endpoint(request=racine.requete_test(), response=Response(),
                    file=fixtures.translation.UploadFile(filename="race.pdf", file=io.BytesIO(fixtures.pdf(3))),
                    target_lang="en", format_options="{}", quality="fast", precise="", pages="",
                    debug="", covered_only="", use_beta="0", x_api_key=fixtures.config.FRONTEND_API_KEY,
                    current_user=account, db=db)
        try:
            results = await asyncio.gather(upload(), upload())
            self.assertEqual(sorted(getattr(result, "status_code", 200) for result in results), [200, 402])
            async with sessions() as db:
                self.assertEqual(await usage.free_pages_used(db, user.id), 3)
        finally:
            async with sessions() as db:
                await db.execute(delete(User).where(User.id == user.id))
                await db.commit()
            await fx.asyncTearDown()
            await engine.dispose()


if __name__ == "__main__":
    unittest.main()
