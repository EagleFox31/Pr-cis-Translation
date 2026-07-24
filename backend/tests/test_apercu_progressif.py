"""L'aperçu se remplit page par page — et il dit la vérité à chaque instant.

CE QUI EST VÉRIFIÉ
  1. SOCLE     — le document d'origine est affichable avant toute traduction,
                 avec TOUTES ses pages.
  2. GREFFE    — une diapositive traduite remplace SA page, et seulement elle ;
                 les autres restent celles de l'original.
  3. ORDRE     — greffer 3 puis 1 ne permute rien : la page 1 reste la page 1.
  4. LOT       — plusieurs diapositives converties d'un coup atterrissent
                 chacune à sa place (c'est le régime sous charge).
  5. ROBUSTESSE— une greffe qui échoue ne perd pas les diapositives du lot et
                 ne laisse jamais un partiel tronqué.
  6. COÛT      — le profil LibreOffice partagé tient sa promesse.

Le montage est SYNTHÉTIQUE : un deck fabriqué ici, et un texte-repère par
diapositive qui permet de dire, en lisant le PDF, quelle version on regarde.

    backend/venv/Scripts/python.exe backend/tests/test_apercu_progressif.py
"""
from __future__ import annotations

import os
import sys
import tempfile
import time

import racine  # noqa: F401  -- met backend/ sur le chemin

import fitz                                              # noqa: E402
from pptx import Presentation                            # noqa: E402
from pptx.util import Emu, Pt                            # noqa: E402

from app.services.progressive_preview import ProgressivePreview   # noqa: E402
from engines import office                               # noqa: E402
from engines.pptx.engine import PPTXTranslatorEngine     # noqa: E402

N_SLIDES = 6
SOURCE = "ORIGINE"
CIBLE = "TRADUIT"


def _deck(chemin: str, marque: str) -> None:
    """Un deck où chaque diapositive porte « <marque> <n> » en gros."""
    prs = Presentation()  # pyrefly: ignore[not-callable]
    for i in range(1, N_SLIDES + 1):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Emu(500000), Emu(1500000),
                                  Emu(8000000), Emu(1500000))
        p = tb.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = f"{marque} {i}"
        r.font.size = Pt(54)
    prs.save(chemin)


def _texte_page(pdf: bytes, index: int) -> str:
    with fitz.open(stream=pdf, filetype="pdf") as d:
        if not (0 <= index < d.page_count):
            return ""
        return d[index].get_text().strip()


def _pages(pdf: bytes) -> int:
    with fitz.open(stream=pdf, filetype="pdf") as d:
        return d.page_count


def main() -> int:
    checks: list[tuple[str, bool, str]] = []

    def ok(nom, cond, detail=""):
        checks.append((nom, bool(cond), detail))

    if not office.SOFFICE_PATH:
        print("LibreOffice introuvable : suite ignorée.")
        return 0

    td = tempfile.mkdtemp(prefix="apercu_")
    origine = os.path.join(td, "origine.pptx")
    traduit = os.path.join(td, "traduit.pptx")
    _deck(origine, SOURCE)
    _deck(traduit, CIBLE)

    # ── 6. Le profil partagé tient-il sa promesse ? ──────────────────────
    with open(origine, "rb") as f:
        brut = f.read()
    t0 = time.time(); socle = office.convert_to_pdf(brut, "pptx"); froid = time.time() - t0
    t0 = time.time(); office.convert_to_pdf(brut, "pptx"); chaud = time.time() - t0
    ok("PROFIL  la 2e conversion est plus rapide que la 1re",
       chaud < froid, f"{froid:.2f}s puis {chaud:.2f}s")
    ok("PROFIL  le gain est SUBSTANTIEL (pas du bruit de mesure)",
       chaud < froid * 0.75, f"{froid:.2f}s -> {chaud:.2f}s")

    # ── 1. Le socle montre tout le document, en langue SOURCE ────────────
    ok("SOCLE  toutes les pages sont là avant toute traduction",
       _pages(socle) == N_SLIDES, f"{_pages(socle)} pages")
    ok("SOCLE  la page 1 est bien l'ORIGINAL",
       SOURCE in _texte_page(socle, 0), _texte_page(socle, 0))

    # Le moteur ouvre le deck TRADUIT : `build_partial_pptx(only_slides=…)` en
    # extrait les diapositives, exactement comme pendant un vrai job.
    eng = PPTXTranslatorEngine()
    eng._extract_zip(traduit)

    # L'objet de PRODUCTION, avec ses deux dependances injectees. Le test
    # exerce ainsi le code livre, et non une reimplementation de sa logique --
    # qui prouverait la technique sans rien prouver de l'implementation.
    apercu = ProgressivePreview(
        extraire=lambda chemin, pages: eng.build_partial_pptx(
            chemin, max(pages), only_slides=pages),
        convertir=lambda octets: office.convert_to_pdf(octets, "pptx"),
        extension="pptx")
    apercu.poser_socle(socle)

    def greffer(nums: list[int]) -> bytes:
        return apercu.greffer(nums)

    # ── 2. Une greffe ne touche QUE sa page ──────────────────────────────
    t0 = time.time()
    p3 = greffer([3])
    latence = time.time() - t0
    ok("GREFFE  le nombre de pages ne bouge pas",
       _pages(p3) == N_SLIDES, f"{_pages(p3)} pages")
    ok("GREFFE  la diapositive 3 est TRADUITE",
       CIBLE in _texte_page(p3, 2), _texte_page(p3, 2))
    ok("GREFFE  la diapositive 1 est restée l'ORIGINAL",
       SOURCE in _texte_page(p3, 0), _texte_page(p3, 0))
    ok("GREFFE  la diapositive 6 est restée l'ORIGINAL",
       SOURCE in _texte_page(p3, 5), _texte_page(p3, 5))
    ok("GREFFE  une page traduite coûte quelques secondes, pas le document",
       latence < 8.0, f"{latence:.2f}s")

    # ── 3. L'ORDRE ne dépend pas de l'ordre d'arrivée ────────────────────
    p31 = greffer([1])
    ok("ORDRE  greffer 3 PUIS 1 laisse la page 1 en tête",
       CIBLE in _texte_page(p31, 0) and "1" in _texte_page(p31, 0),
       _texte_page(p31, 0))
    ok("ORDRE  la page 3 greffée avant n'a pas bougé",
       CIBLE in _texte_page(p31, 2) and "3" in _texte_page(p31, 2),
       _texte_page(p31, 2))
    ok("ORDRE  la page 2, jamais greffée, est toujours l'ORIGINAL",
       SOURCE in _texte_page(p31, 1) and "2" in _texte_page(p31, 1),
       _texte_page(p31, 1))

    # ── 4. Un LOT (régime sous charge) place chaque page correctement ────
    t0 = time.time()
    plot = greffer([2, 5, 6])
    t_lot = time.time() - t0
    for num in (2, 5, 6):
        ok(f"LOT  la diapositive {num} est traduite ET à sa place",
           CIBLE in _texte_page(plot, num - 1)
           and str(num) in _texte_page(plot, num - 1),
           _texte_page(plot, num - 1))
    ok("LOT  la diapositive 4, jamais demandée, reste l'ORIGINAL",
       SOURCE in _texte_page(plot, 3), _texte_page(plot, 3))
    # C'est TOUT l'intérêt de la coalescence : trois pages pour bien moins que
    # trois fois le prix d'une.
    ok("LOT  trois pages d'un coup coûtent moins que trois greffes séparées",
       t_lot < latence * 2.5, f"lot de 3 : {t_lot:.2f}s / une seule : {latence:.2f}s")

    # ── 5. Le partiel n'est jamais tronqué ───────────────────────────────
    ok("ROBUSTESSE  le partiel garde toujours toutes ses pages",
       all(_pages(p) == N_SLIDES for p in (p3, p31, plot)),
       f"{[_pages(p) for p in (p3, p31, plot)]}")

    # ── 5-bis. Un appariement INCERTAIN doit lever, pas deviner ──────────
    # Si le lot rendu ne compte pas autant de pages que demandé, on ne sait
    # plus quelle page correspond a quelle diapositive. Greffer quand meme,
    # c'est afficher une traduction en face de la mauvaise page -- un mensonge
    # silencieux. On exige donc une exception, et un apercu INTACT.
    # On compare le CONTENU, pas les octets : la serialisation d'un PDF n'est
    # pas deterministe (identifiants, table des objets), et un test sur les
    # octets echouerait sans qu'aucun pixel n'ait change.
    def _empreinte(pdf: bytes):
        return [_texte_page(pdf, i) for i in range(_pages(pdf))]

    avant = _empreinte(apercu.assembler())
    n_avant = apercu.pages_traduites
    leve = False
    try:
        apercu.greffer([N_SLIDES + 40])      # diapositive qui n'existe pas
    except Exception:
        leve = True
    ok("APPARIEMENT  un lot qui ne correspond pas leve une exception", leve)
    ok("APPARIEMENT  aucune page n'est greffee a l'aveugle",
       apercu.pages_traduites == n_avant,
       f"{n_avant} -> {apercu.pages_traduites}")
    ok("APPARIEMENT  l'apercu affiche EXACTEMENT ce qu'il affichait",
       _empreinte(apercu.assembler()) == avant)

    eng._cleanup_temp()

    print()
    n_ok = sum(1 for _, c, _ in checks if c)
    for nom, cond, detail in checks:
        ligne = ("  OK  " if cond else " FAIL ") + f"  {nom}"
        if not cond and detail:
            ligne += f"   [{detail}]"
        print(ligne)
    print(f"\n{n_ok}/{len(checks)}")
    return 0 if n_ok == len(checks) else 1


if __name__ == "__main__":
    sys.exit(main())
