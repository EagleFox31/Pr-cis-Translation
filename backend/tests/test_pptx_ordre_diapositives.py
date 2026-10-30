"""PPTX : la diapositive n°N est celle que `presentation.xml` affiche en N.

LE DÉFAUT QUE CETTE SUITE VERROUILLE
------------------------------------
Le moteur supposait que la diapositive n°1 était `slide1.xml`. Un export
Google Slides nomme `slide1000.xml`, `slide1001.xml`… : le moteur comptait
bien 38 fichiers, cherchait `slide1.xml` à `slide38.xml`, n'en trouvait aucun
et rangeait chaque diapositive dans « rien à traduire ». Le job passait en
succès 38/38 sans un mot traduit, et le PPTX final ne gardait aucune
diapositive (sélection « numéro de fichier ≤ 38 »). Réordonner des
diapositives dans PowerPoint (sans renommer) donnait, lui, un aperçu dans le
désordre.

Le deck de test est SYNTHÉTIQUE : fichiers renommés `slide1001.xml…` ET ordre
d'affichage inversé, pour que ni le nom ni le tri numérique ne tombent juste
par hasard.

    backend/venv/Scripts/python.exe backend/tests/test_pptx_ordre_diapositives.py
"""
from __future__ import annotations

import json
import os
import re
import shutil
import sys
import tempfile
import zipfile

import racine  # noqa: F401  -- met backend/ sur le chemin

from lxml import etree                                    # noqa: E402
from pptx import Presentation                             # noqa: E402
from pptx.util import Inches                              # noqa: E402

from engines.pptx.engine import PPTXTranslatorEngine, NAMESPACES  # noqa: E402
from app.services import translation_runner as tr         # noqa: E402
from app.services.jobs import jobs                        # noqa: E402

N = 4


def _deck_google(path: str, textes: list[str] | None = None) -> None:
    """Deck de N diapositives, fichiers `slide100k.xml`, ordre d'affichage
    inversé : la diapositive affichée en 1 porte « TEXTE N »."""
    brut = path + ".brut.pptx"
    prs = Presentation()  # pyrefly: ignore[not-callable]
    for i in range(1, N + 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if textes is None or textes[i - 1]:
            zone = slide.shapes.add_textbox(Inches(1), Inches(1),
                                            Inches(6), Inches(1))
            zone.text_frame.text = (textes[i - 1] if textes else f"TEXTE {i}")
    prs.save(brut)

    def renomme(nom: str) -> str:
        return re.sub(r"slides/(_rels/)?slide(\d+)\.xml",
                      lambda m: f"slides/{m.group(1) or ''}slide{1000 + int(m.group(2))}.xml",
                      nom)

    with zipfile.ZipFile(brut) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename in ("ppt/_rels/presentation.xml.rels", "[Content_Types].xml"):
                data = re.sub(rb"slide(\d+)\.xml",
                              lambda m: b"slide%d.xml" % (1000 + int(m.group(1))), data)
            elif info.filename == "ppt/presentation.xml":
                root = etree.fromstring(data)
                lst = root.find("p:sldIdLst", NAMESPACES)
                ids = list(lst)
                for el in ids:
                    lst.remove(el)
                for el in reversed(ids):
                    lst.append(el)
                data = etree.tostring(root, encoding="utf-8", xml_declaration=True,
                                      standalone=True)
            zout.writestr(renomme(info.filename), data)
    os.remove(brut)


def _textes(path: str) -> list[str]:
    """Texte de chaque diapositive, dans l'ordre d'AFFICHAGE."""
    prs = Presentation(path)  # pyrefly: ignore[not-callable]
    return [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
            for s in prs.slides]


def _carte(slide_data: dict, prefixe: str) -> dict:
    return {el["id"]: _prefixe(el["text"], prefixe) for el in slide_data["text_elements"]}


def _prefixe(texte: str, prefixe: str) -> str:
    """Préfixe le texte APRÈS le marqueur de run `[[n]]` éventuel."""
    return texte.replace("]]", "]]" + prefixe, 1) if "]]" in texte else prefixe + texte


class Traducteur:
    """Préfixe « EN: » chaque fragment, sans réseau."""

    def fiche_document(self, *_a, **_k):
        return ""

    def translate_json(self, chemin, target_lang, progress_callback=None, **_):
        with open(chemin, encoding="utf-8") as f:
            data = json.load(f)
        for sd in data["slides"]:
            for el in tr._elements_de(sd):
                el["translated_text"] = el["text"].replace("]]", "]]EN:", 1)
        sortie = chemin + ".out.json"
        with open(sortie, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False)
        return True, sortie

    def _translate_batch(self, *a, **k):
        return None


def _lancer(dossier: str, octets: bytes, nom: str) -> tuple[str, str, dict]:
    job_id = jobs.create()
    sortie = os.path.join(dossier, f"render_{nom}.pptx")
    traduction = os.path.join(dossier, f"translated_{nom}.json")
    tr.run_pptx_progressive_job(
        job_id, octets, os.path.join(dossier, f"original_{nom}.pptx"), sortie,
        f"{nom}.pptx", os.path.join(dossier, f"partial_{nom}.pdf"), traduction, "en")
    return sortie, traduction, jobs.get(job_id) or {}


def main() -> int:
    checks: list[tuple[str, bool, str]] = []

    def ok(nom, cond, detail=""):
        checks.append((nom, bool(cond), detail))

    dossier = tempfile.mkdtemp(prefix="ordre_pptx_")
    try:
        deck = os.path.join(dossier, "google.pptx")
        _deck_google(deck)
        attendu = [f"TEXTE {i}" for i in range(N, 0, -1)]
        ok("0. fixture : ordre d'affichage inversé", _textes(deck) == attendu,
           str(_textes(deck)))

        # ── 1. Pipeline progressif : extraction / injection par POSITION ──
        eng = PPTXTranslatorEngine()
        try:
            eng._cleanup_temp()
            eng._extract_zip(deck)
            ok("1. slide_count", eng.slide_count() == N, str(eng.slide_count()))
            lus = []
            for pos in range(1, N + 1):
                sd, info = eng.extract_slide(pos)
                lus.append(sd and " ".join(el["text"] for el in sd["text_elements"]))
                if sd:
                    eng.inject_slide(pos, _carte(sd, "EN:"))
            ok("1. extraction dans l'ordre d'affichage",
               all(lus) and all(f"TEXTE {N + 1 - p}" in lus[p - 1] for p in range(1, N + 1)),
               str(lus))
            try:
                eng.extract_slide(N + 1)
                ok("1. diapositive introuvable : erreur", False)
            except ValueError:
                ok("1. diapositive introuvable : erreur", True)

            final = os.path.join(dossier, "final.pptx")
            eng.build_partial_pptx(final, N)
            t = _textes(final)
            ok("1. PPTX final : toutes les diapositives, traduites, dans l'ordre",
               t == [f"EN:{x}" for x in attendu], str(t))

            isolee = os.path.join(dossier, "isolee.pptx")
            eng.build_partial_pptx(isolee, 2, only_slides={2})
            t = _textes(isolee)
            ok("1. aperçu : la position 2 isolée est la 2e affichée",
               t == [f"EN:{attendu[1]}"], str(t))
        finally:
            eng._cleanup_temp()

        # ── 2. Rendu à la demande : JSON progressif → inject_translation ──
        # L'aperçu rouvert depuis la bibliothèque passe par ce chemin : les
        # identifiants doivent concorder avec ceux du pipeline progressif.
        eng = PPTXTranslatorEngine()
        try:
            data, _ = eng.extract_text(deck, os.path.join(dossier, "ext.json"))
            ok("2. extract_text : N diapositives", data and len(data["slides"]) == N)
            for sd in (data or {"slides": []})["slides"]:
                for el in sd["text_elements"]:
                    el["translated_text"] = _prefixe(el["text"], "EN:")
            jpath = os.path.join(dossier, "trad.json")
            with open(jpath, "w", encoding="utf-8") as f:
                json.dump(data, f)
            rendu = os.path.join(dossier, "rendu.pptx")
            okinj, msg = eng.inject_translation(deck, jpath, rendu)
            t = _textes(rendu) if okinj else []
            ok("2. inject_translation : traduit, dans l'ordre",
               t == [f"EN:{x}" for x in attendu], f"{msg} {t}")
        finally:
            eng._cleanup_temp()

        # ── 3. Job réel sur le deck renommé ──────────────────────────────
        tr.convert_to_pdf_bytes = lambda *a, **k: (_ for _ in ()).throw(
            RuntimeError("LibreOffice neutralisé"))
        tr.ai_translator, tr.ai_active = Traducteur(), True
        with open(deck, "rb") as f:
            octets = f.read()
        sortie, traduction, job = _lancer(dossier, octets, "job")
        ok("3. job terminé", job.get("state") == "done", str(job.get("error")))
        if os.path.isfile(traduction):
            with open(traduction, encoding="utf-8") as f:
                n = len(json.load(f)["slides"])
        else:
            n = 0
        ok("3. traduction enregistrée : N diapositives", n == N, str(n))
        t = _textes(sortie) if os.path.isfile(sortie) else []
        ok("3. PPTX final traduit, N diapositives dans l'ordre",
           len(t) == N and all("EN:" in x and attendu[i] in x for i, x in enumerate(t)),
           str(t))

        # ── 4. Garde-fou : aucune diapositive ne livre de texte → échec ───
        # Le gabarit python-pptx porte du texte de mise en page : on force
        # donc l'extraction à ne rien rendre, comme le faisait le défaut.
        extraire = PPTXTranslatorEngine.extract_slide
        PPTXTranslatorEngine.extract_slide = (
            lambda self, n, *a, **k: (None, "Aucun texte trouvé dans cette slide."))
        try:
            _, _, job = _lancer(dossier, octets, "vide")
        finally:
            PPTXTranslatorEngine.extract_slide = extraire
        ok("4. aucun texte extrait : job en erreur, message explicite",
           job.get("state") == "error" and "Aucun texte" in (job.get("error") or ""),
           f"{job.get('state')} {job.get('error')}")
    finally:
        shutil.rmtree(dossier, ignore_errors=True)

    for nom, passe, detail in checks:
        print(("OK   " if passe else "ECHEC"), nom, ("" if passe else detail))
    n_ok = sum(1 for _, p, _ in checks if p)
    print(f"{n_ok}/{len(checks)}")
    return 0 if n_ok == len(checks) else 1


if __name__ == "__main__":
    sys.exit(main())
