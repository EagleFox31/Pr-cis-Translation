"""Beta lifecycle, API authorization and translation admission; no AI/payment calls."""
from __future__ import annotations
import asyncio
from datetime import timedelta
import io
import json
import os
from pathlib import Path
import secrets
import tempfile
import unittest
from unittest.mock import patch
import uuid
import zipfile
from urllib.parse import urlsplit

import racine  # noqa: F401

os.environ.setdefault("PRECIS_NO_BANNER", "1")
os.environ.setdefault("EMAIL_ENABLED", "false")
os.environ.setdefault("FRONTEND_API_KEY", secrets.token_hex(24))
os.environ.setdefault("JWT_SECRET", secrets.token_hex(32))
os.environ.setdefault("DEEPSEEK_API_KEY", secrets.token_hex(24))

import fitz
from docx import Document as WordDocument
from pptx import Presentation
import httpx
from sqlalchemy.exc import IntegrityError
from sqlalchemy import delete, event, select
from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker
from sqlalchemy.pool import NullPool
from app import create_app, config
from app.api import translate as translation
from app.core.database import get_db, DATABASE_URL
from app.core.security import create_access_token
from app.models import Base, User, Document, UsageRecord
from app.models.beta import BetaCampaign, BetaEnrollment, BetaUsage
from app.services import beta, usage
from app.services.beta_meter import BetaCost, DEFAULT_LIMITS, office_cost, word_count
from app.services.jobs import jobs
from app.rate_limit import limiter


def pdf(pages=1):
    doc = fitz.open()
    for number in range(pages):
        page = doc.new_page()
        page.insert_text((40, 80), f"Beta fixture page {number} {uuid.uuid4().hex}")
    data = doc.tobytes()
    doc.close()
    return data


def xlsx(words=4, sheets=1):
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        entries = "".join(f'<sheet name="Table{n}" sheetId="{n}" r:id="rId{n}"/>' for n in range(1, sheets + 1))
        archive.writestr("xl/workbook.xml", f'<workbook xmlns="{ns}" xmlns:r="{rns}"><sheets>{entries}</sheets></workbook>')
        for n in range(1, sheets + 1):
            text = " ".join(["bonjour"] * words) if n == 1 else "texte"
            archive.writestr(f"xl/worksheets/sheet{n}.xml",
                f'<worksheet xmlns="{ns}"><sheetData><row r="1">'
                f'<c r="A1" t="inlineStr"><is><t>{text}</t></is></c>'
                '<c r="B1"><v>123</v></c><c r="C1"><f>SUM(B1)</f><v>123</v></c>'
                '</row></sheetData></worksheet>')
    return out.getvalue()


class MeterTests(unittest.TestCase):
    def test_word_meter_uses_real_text_and_ignores_style_markers(self):
        document = WordDocument()
        document.add_paragraph("bonjour " * 620)
        document.sections[0].header.paragraphs[0].text = "entete document"
        data = io.BytesIO(); document.save(data)
        cost = office_cost(data.getvalue(), "docx")
        self.assertEqual(cost.units, 622)
        self.assertEqual(cost.format, "docx")
        self.assertEqual(word_count("[[0]]hello world[[/0]] 123"), 2)
        self.assertEqual(word_count("中文"), 2)
        # Unspaced Thai counts by length (25 characters = 5 words), not as one word.
        self.assertEqual(word_count("กขคงจฉชซฌญฎฏฐฑฒณดตถทธนบปผ hello"), 6)
        self.assertEqual(word_count("สวัสดีครับ"), 2)

    def test_excel_sheets_and_all_extracted_text_are_measured(self):
        cost = office_cost(xlsx(words=20, sheets=3), "xlsx")
        self.assertEqual(cost.units, 3)
        self.assertGreaterEqual(cost.words, 25)  # cells plus the translated tab names
        huge = office_cost(xlsx(words=5001), "xlsx")
        self.assertGreater(huge.words, 5000)

    def test_invalid_office_container_does_not_get_a_one_page_allowance(self):
        with self.assertRaises(ValueError):
            office_cost(b"not an Office file", "xlsx")


class BetaTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="precis-beta-test-")
        self.engine = create_async_engine("sqlite+aiosqlite://")
        @event.listens_for(self.engine.sync_engine, "connect")
        def foreign_keys(connection, _record):
            connection.execute("PRAGMA foreign_keys=ON")
        async with self.engine.begin() as db:
            await db.run_sync(Base.metadata.create_all)
        self.sessions = async_sessionmaker(self.engine, expire_on_commit=False)
        self.app = create_app()
        if limiter:
            limiter.reset()
        async def database():
            async with self.sessions() as db:
                yield db
        self.app.dependency_overrides[get_db] = database
        self.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=self.app), base_url="http://test")
        async with self.sessions() as db:
            self.admin = User(email="beta-admin@example.com", plan="admin", email_verified=True)
            self.tester = User(email="beta-tester@example.com", plan="free", email_verified=True, page_credits=7,
                               created_at=racine.compte_ancien())
            self.other = User(email="beta-other@example.com", plan="free", email_verified=True,
                              created_at=racine.compte_ancien())
            db.add_all([self.admin, self.tester, self.other]); await db.commit()
        self.admin_headers = self.headers(self.admin)
        self.tester_headers = self.headers(self.tester)
        self.path_patch = patch.object(translation, "build_job_paths", side_effect=lambda digest, language: (
            str(Path(self.temp.name) / digest), str(Path(self.temp.name) / digest / language)))
        self.path_patch.start()
        self.queued = []
        self.scheduler_patch = patch.object(translation.scheduler, "submit",
                                           side_effect=lambda job_id, *_: self.queued.append(job_id))
        self.scheduler_patch.start()

    def headers(self, user):
        return {"Authorization": "Bearer " + create_access_token(user.id, user.email),
                "X-API-Key": config.FRONTEND_API_KEY}

    async def asyncTearDown(self):
        for job_id in self.queued:
            jobs.forget(job_id)
        self.scheduler_patch.stop(); self.path_patch.stop()
        await self.client.aclose(); await self.engine.dispose(); self.temp.cleanup()

    async def start_campaign(self, **overrides):
        payload = {"name": "Private beta", "capacity": 40, "first_wave_capacity": 20,
                   "access_days": 30, "limits": DEFAULT_LIMITS, **overrides}
        result = await self.client.post("/api/admin/beta/campaigns", json=payload, headers=self.admin_headers)
        self.assertEqual(result.status_code, 201, result.text)
        self.assertFalse(result.json()["recruitment_open"])
        self.campaign_id = result.json()["id"]
        opened = await self.client.patch(f"/api/admin/beta/campaigns/{self.campaign_id}", json={**payload, "state": "open"},
                                         headers=self.admin_headers)
        self.assertEqual(opened.status_code, 200, opened.text)
        return payload

    async def invite_tester(self, user=None):
        user = user or self.tester
        result = await self.client.post("/api/admin/beta/invitations",
            json={"email": user.email.upper(), "primary_format": "xlsx"}, headers=self.admin_headers)
        self.assertEqual(result.status_code, 201, result.text)
        return result.json()

    async def activate_tester(self):
        await self.start_campaign(); await self.invite_tester()
        result = await self.client.post("/api/beta/activate", headers=self.tester_headers)
        self.assertEqual(result.status_code, 200, result.text)
        return result.json()

    async def translate(self, data, filename="file.pdf", **fields):
        return await self.client.post("/api/translate", headers=self.tester_headers,
            files={"file": (filename, data)}, data={"target_lang": "en", "use_beta": "1", **fields})

    async def test_admin_authorization_and_email_bound_activation(self):
        self.assertEqual((await self.client.get("/api/admin/beta")).status_code, 401)
        self.assertEqual((await self.client.get("/api/admin/beta", headers=self.tester_headers)).status_code, 403)
        await self.start_campaign(); invite = await self.invite_tester()
        self.assertEqual(invite["email"], self.tester.email)
        other = await self.client.post("/api/beta/activate", headers=self.headers(self.other))
        self.assertEqual(other.status_code, 404)
        async with self.sessions() as db:
            (await db.get(User, self.tester.id)).email_verified = False; await db.commit()
        self.assertEqual((await self.client.post("/api/beta/activate", headers=self.tester_headers)).status_code, 403)

    async def test_invitation_activates_on_status_read(self):
        await self.start_campaign(); await self.invite_tester()
        async with self.sessions() as db:
            (await db.get(User, self.tester.id)).email_verified = False; await db.commit()
        pending = (await self.client.get("/api/beta/me", headers=self.tester_headers)).json()
        self.assertEqual(pending["enrollment"]["state"], "invited")
        async with self.sessions() as db:
            (await db.get(User, self.tester.id)).email_verified = True; await db.commit()
        first = (await self.client.get("/api/beta/me", headers=self.tester_headers)).json()
        self.assertEqual(first["enrollment"]["state"], "active")
        second = (await self.client.get("/api/beta/me", headers=self.tester_headers)).json()
        self.assertEqual(first["enrollment"]["expires_at"], second["enrollment"]["expires_at"])
        other = (await self.client.get("/api/beta/me", headers=self.headers(self.other))).json()
        self.assertIsNone(other["enrollment"])

    async def test_two_waves_and_two_month_recruitment_window(self):
        await self.start_campaign(capacity=2, first_wave_capacity=1)
        await self.invite_tester()
        second = await self.client.post("/api/admin/beta/invitations",
            json={"email": self.other.email}, headers=self.admin_headers)
        self.assertEqual(second.status_code, 409)
        async with self.sessions() as db:
            campaign = await db.get(BetaCampaign, self.campaign_id)
            campaign.starts_at = beta.now_utc() - timedelta(days=32)
            campaign.recruitment_ends_at = beta.now_utc() + timedelta(days=27)
            await db.commit()
        self.assertEqual((await self.client.post("/api/admin/beta/invitations",
            json={"email": self.other.email}, headers=self.admin_headers)).status_code, 201)
        async with self.sessions() as db:
            campaign = await db.get(BetaCampaign, self.campaign_id)
            campaign.recruitment_ends_at = beta.now_utc() - timedelta(seconds=1); await db.commit()
        self.assertEqual((await self.client.post("/api/beta/activate", headers=self.tester_headers)).status_code, 409)

    async def test_non_beta_format_ignores_beta_choice(self):
        await self.activate_tester()
        result = await self.translate(b"Bonjour le monde", filename="note.txt")
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(len(self.queued), 1)

    async def test_activation_is_once_and_expiry_cannot_spend_purchased_credits(self):
        first = await self.activate_tester()
        second = (await self.client.post("/api/beta/activate", headers=self.tester_headers)).json()
        self.assertEqual(first["expires_at"], second["expires_at"])
        async with self.sessions() as db:
            enrollment = await db.get(BetaEnrollment, first["id"])
            enrollment.expires_at = beta.now_utc() - timedelta(seconds=1); await db.commit()
        third = (await self.client.post("/api/beta/activate", headers=self.tester_headers)).json()
        self.assertEqual(third["state"], "expired")
        result = await self.translate(pdf())
        self.assertEqual(result.status_code, 402)
        self.assertEqual(result.json()["reason"], "beta_inactive")
        async with self.sessions() as db:
            self.assertEqual((await db.get(User, self.tester.id)).page_credits, 7)
        self.assertEqual(self.queued, [])

    async def test_beta_pdf_is_clear_and_separate_from_subscription_and_paid_credits(self):
        enrollment = await self.activate_tester()
        result = await self.translate(pdf(3))
        self.assertEqual(result.status_code, 200, result.text)
        async with self.sessions() as db:
            document = (await db.execute(select(Document))).scalar_one()
            user = await db.get(User, self.tester.id)
            self.assertTrue(document.paid)
            self.assertEqual(user.page_credits, 7)
            self.assertEqual(user.plan, "free")
            self.assertEqual((await db.execute(select(UsageRecord))).scalars().all(), [])
            self.assertEqual((await beta.totals(db, enrollment["id"]))["pdf"]["units"], 3)
            self.assertNotIn("xlsx", await beta.totals(db, enrollment["id"]))

    async def test_each_office_format_consumes_its_own_allowance(self):
        enrollment = await self.activate_tester()
        doc = WordDocument(); doc.add_paragraph("bonjour " * 40)
        stream = io.BytesIO(); doc.save(stream)
        self.assertEqual((await self.translate(stream.getvalue(), filename="file.docx")).status_code, 200)
        deck = Presentation()
        for _ in range(2):
            deck.slides.add_slide(deck.slide_layouts[6])
        stream = io.BytesIO(); deck.save(stream)
        self.assertEqual((await self.translate(stream.getvalue(), filename="file.pptx")).status_code, 200)
        self.assertEqual((await self.translate(xlsx(5, 2), filename="file.xlsx")).status_code, 200)
        async with self.sessions() as db:
            used = await beta.totals(db, enrollment["id"])
            self.assertEqual(used["docx"]["units"], 40)
            self.assertEqual(used["pptx"]["units"], 2)
            self.assertEqual(used["xlsx"]["units"], 2)
            self.assertNotIn("pdf", used)
            self.assertEqual((await db.get(User, self.tester.id)).page_credits, 7)

    async def test_beta_scan_recognizes_all_covered_pages_not_just_the_free_trial_page(self):
        enrollment = await self.activate_tester()
        scan = fitz.open()
        for _ in range(3):
            scan.new_page()
        data = scan.tobytes(); scan.close()
        with patch.object(translation, "prepare_scanned_pdf", return_value=(pdf(3), 3)) as recognize:
            result = await self.translate(data)
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(recognize.call_args.args[1], {1, 2, 3})
        async with self.sessions() as db:
            self.assertEqual((await beta.totals(db, enrollment["id"]))["pdf"]["units"], 3)

    async def test_excel_dual_cap_and_word_quota_apply_before_a_job_is_queued(self):
        await self.activate_tester()
        too_many_words = await self.translate(xlsx(5001), filename="budget.xlsx")
        self.assertEqual(too_many_words.status_code, 402, too_many_words.text)
        self.assertEqual(too_many_words.json()["reason"], "beta_quota")
        too_many_sheets = await self.translate(xlsx(1, 11), filename="budget.xlsx")
        self.assertEqual(too_many_sheets.status_code, 402)
        document = WordDocument(); document.add_paragraph("bonjour " * 10001)
        stream = io.BytesIO(); document.save(stream)
        self.assertEqual((await self.translate(stream.getvalue(), filename="budget.docx")).status_code, 402)
        self.assertEqual(self.queued, [])
        self.assertEqual((await self.translate(xlsx(10, 2), filename="small.xlsx")).status_code, 200)

    async def test_refund_rebill_and_document_deletion_preserve_the_correct_usage(self):
        enrollment = await self.activate_tester()
        self.assertEqual((await self.translate(pdf(2))).status_code, 200)
        async with self.sessions() as db:
            document = (await db.execute(select(Document))).scalar_one()
            document_id = document.id
            await usage.refund_for_document(db, document_id)
            await usage.refund_for_document(db, document_id); await db.commit()
            self.assertEqual(await beta.totals(db, enrollment["id"]), {})
            await usage.rebill_for_document(db, document_id)
            await usage.rebill_for_document(db, document_id); await db.commit()
            self.assertEqual((await beta.totals(db, enrollment["id"]))["pdf"]["units"], 2)
            await db.delete(document); await db.commit()
            rows = (await db.execute(select(BetaUsage))).scalars().all()
            self.assertEqual(len(rows), 1); self.assertIsNone(rows[0].document_id)
            self.assertEqual((await beta.totals(db, enrollment["id"]))["pdf"]["units"], 2)
            # A later worker failure can still refund a document deleted while
            # it was running; the stable ledger key survives the SET NULL FK.
            await usage.refund_for_document(db, document_id); await db.commit()
            self.assertEqual(await beta.totals(db, enrollment["id"]), {})

    async def test_revocation_and_restore_do_not_reset_time_or_quotas(self):
        enrollment = await self.activate_tester()
        await self.translate(pdf(2))
        path = "/api/admin/beta/enrollments/" + enrollment["id"]
        revoked = await self.client.patch(path, json={"revoke": True}, headers=self.admin_headers)
        self.assertEqual(revoked.json()["state"], "revoked")
        self.assertEqual((await self.translate(pdf())).status_code, 402)
        restored = await self.client.patch(path, json={"revoke": False}, headers=self.admin_headers)
        self.assertEqual(restored.json()["expires_at"], enrollment["expires_at"])
        self.assertEqual(restored.json()["formats"]["pdf"]["used"], 2)

    async def test_explicit_regular_plan_choice_uses_free_quota_without_beta_usage(self):
        enrollment = await self.activate_tester()
        result = await self.translate(pdf(2), use_beta="0")
        self.assertEqual(result.status_code, 200, result.text)
        async with self.sessions() as db:
            self.assertEqual(await beta.totals(db, enrollment["id"]), {})
            self.assertEqual((await db.get(User, self.tester.id)).page_credits, 7)
            self.assertEqual(await usage.free_pages_used(db, self.tester.id), 2)

    async def test_default_updates_do_not_change_existing_allowances_or_recruitment_dates(self):
        enrollment = await self.activate_tester()
        before = (await self.client.get("/api/admin/beta", headers=self.admin_headers)).json()["campaign"]
        changed = await self.client.patch(f"/api/admin/beta/campaigns/{self.campaign_id}",
            json={"name": "Renamed program"}, headers=self.admin_headers)
        self.assertEqual(changed.status_code, 200, changed.text)
        self.assertEqual(changed.json()["starts_at"], before["starts_at"])
        current = (await self.client.get("/api/beta/me", headers=self.tester_headers)).json()["enrollment"]
        self.assertEqual(current["limits"], enrollment["limits"])
        self.assertEqual(current["expires_at"], enrollment["expires_at"])


    async def open_second_program(self, **overrides):
        first = self.campaign_id
        closed = await self.client.patch(f"/api/admin/beta/campaigns/{first}", json={"state": "closed"},
                                         headers=self.admin_headers)
        self.assertEqual(closed.status_code, 200, closed.text)
        listing = (await self.client.get("/api/admin/beta/campaigns", headers=self.admin_headers)).json()
        created = await self.client.post("/api/admin/beta/campaigns",
            json={**listing["defaults"], "name": "Second program", **overrides}, headers=self.admin_headers)
        self.assertEqual(created.status_code, 201, created.text)
        self.campaign_id = created.json()["id"]
        opened = await self.client.patch(f"/api/admin/beta/campaigns/{self.campaign_id}", json={"state": "open"},
                                         headers=self.admin_headers)
        self.assertEqual(opened.status_code, 200, opened.text)
        return first

    async def test_programs_are_listed_prefilled_and_only_one_is_open(self):
        await self.start_campaign(name="First program", capacity=12, first_wave_capacity=6, access_days=21)
        first = self.campaign_id
        listing = await self.client.get("/api/admin/beta/campaigns", headers=self.admin_headers)
        self.assertEqual(listing.status_code, 200, listing.text)
        self.assertEqual([row["id"] for row in listing.json()["campaigns"]], [first])
        defaults = listing.json()["defaults"]
        self.assertEqual((defaults["name"], defaults["capacity"], defaults["access_days"]), ("First program", 12, 21))
        draft = await self.client.post("/api/admin/beta/campaigns", json={**defaults, "name": "Next program"},
                                       headers=self.admin_headers)
        self.assertEqual(draft.status_code, 201, draft.text)
        self.assertEqual(draft.json()["state"], "draft")
        refused = await self.client.patch(f"/api/admin/beta/campaigns/{draft.json()['id']}",
                                          json={"state": "open"}, headers=self.admin_headers)
        self.assertEqual(refused.status_code, 409)
        self.assertIn("campaign_already_open", refused.text)
        current = (await self.client.get("/api/admin/beta", headers=self.admin_headers)).json()
        self.assertEqual(current["campaign"]["id"], first, "the open program stays the current one")
        self.assertEqual(len(current["campaigns"]), 2)
        chosen = await self.client.get("/api/admin/beta", params={"campaign_id": draft.json()["id"]},
                                       headers=self.admin_headers)
        self.assertEqual(chosen.json()["campaign"]["name"], "Next program")
        self.assertEqual(chosen.json()["enrollments"], [])
        missing = await self.client.get("/api/admin/beta", params={"campaign_id": "nope"}, headers=self.admin_headers)
        self.assertEqual(missing.status_code, 404)
        async with self.sessions() as db:
            db.add(BetaCampaign(name="Racing", state="open", limits=DEFAULT_LIMITS))
            with self.assertRaises(IntegrityError):
                await db.commit()

    async def test_former_tester_joins_the_next_program_once_their_access_ended(self):
        enrollment = await self.activate_tester()
        await self.open_second_program()
        application = {"primary_format": "pdf", "motivation": "Rapports trimestriels bilingues"}
        early = await self.client.post("/api/beta/applications", json=application, headers=self.tester_headers)
        self.assertEqual(early.status_code, 409)
        self.assertIn("previous_access_active", early.text)
        invited = await self.client.post("/api/admin/beta/invitations", json={"email": self.tester.email},
                                         headers=self.admin_headers)
        self.assertEqual(invited.status_code, 409)
        self.assertIn("previous_access_active", invited.text)
        me = (await self.client.get("/api/beta/me", headers=self.tester_headers)).json()
        self.assertEqual(me["enrollment"]["id"], enrollment["id"], "the earlier access keeps running")
        async with self.sessions() as db:
            (await db.get(BetaEnrollment, enrollment["id"])).expires_at = beta.now_utc() - timedelta(seconds=1)
            await db.commit()
        applied = await self.client.post("/api/beta/applications", json=application, headers=self.tester_headers)
        self.assertEqual(applied.status_code, 200, applied.text)
        approved = await self.client.patch(f"/api/admin/beta/applications/{applied.json()['id']}",
                                           json={"decision": "approved"}, headers=self.admin_headers)
        self.assertEqual(approved.status_code, 200, approved.text)
        activated = await self.client.post("/api/beta/activate", headers=self.tester_headers)
        self.assertEqual(activated.status_code, 200, activated.text)
        self.assertNotEqual(activated.json()["id"], enrollment["id"])
        self.assertEqual(activated.json()["state"], "active")
        dashboard = (await self.client.get("/api/admin/beta", params={"campaign_id": self.campaign_id},
                                           headers=self.admin_headers)).json()
        self.assertEqual([row["id"] for row in dashboard["enrollments"]], [activated.json()["id"]])
        me = (await self.client.get("/api/beta/me", headers=self.tester_headers)).json()
        self.assertEqual((me["enrollment"]["id"], me["campaign_name"]), (activated.json()["id"], "Second program"))

    async def test_finished_access_gives_way_to_the_open_program_form(self):
        enrollment = await self.activate_tester()
        async with self.sessions() as db:
            (await db.get(BetaEnrollment, enrollment["id"])).expires_at = beta.now_utc() - timedelta(seconds=1)
            await db.commit()
        await self.open_second_program(access_days=14)
        me = (await self.client.get("/api/beta/me", headers=self.tester_headers)).json()
        self.assertIsNone(me["enrollment"], "an expired earlier access must not hide the application form")
        self.assertEqual(me["access_days"], 14)
        # A pending invitation of a closed program reports that program's own settings.
        await self.client.patch(f"/api/admin/beta/campaigns/{self.campaign_id}", json={"state": "closed"},
                                headers=self.admin_headers)
        async with self.sessions() as db:
            db.add(BetaEnrollment(campaign_id=self.campaign_id, email=self.other.email, primary_format="pdf",
                                  limits=DEFAULT_LIMITS))
            db.add(BetaCampaign(name="Draft", state="draft", access_days=7, limits=DEFAULT_LIMITS))
            await db.commit()
        other = (await self.client.get("/api/beta/me", headers=self.headers(self.other))).json()
        self.assertEqual((other["campaign_name"], other["access_days"], other["can_activate"]),
                         ("Second program", 14, False))

    async def test_an_alias_mailbox_cannot_be_invited_while_its_access_runs(self):
        await self.activate_tester()
        await self.open_second_program()
        async with self.sessions() as db:
            (await db.get(User, self.tester.id)).email = "beta-tester@outlook.com"
            enrollment = (await db.execute(select(BetaEnrollment))).scalars().one()
            enrollment.email = "beta-tester@outlook.com"; enrollment.mailbox_key = "beta-tester@outlook.com"
            await db.commit()
        invited = await self.client.post("/api/admin/beta/invitations", headers=self.admin_headers,
            json={"email": "beta-tester+2@outlook.com", "campaign_id": self.campaign_id})
        self.assertEqual(invited.status_code, 409)
        self.assertIn("previous_access_active", invited.text)

    async def test_an_alias_mailbox_cannot_hold_a_second_running_access(self):
        await self.activate_tester()
        await self.open_second_program()
        async with self.sessions() as db:
            alias = User(email="beta-tester+2@outlook.com", plan="free", email_verified=True,
                         created_at=racine.compte_ancien())
            (await db.get(User, self.tester.id)).email = "beta-tester@outlook.com"
            enrollment = (await db.execute(select(BetaEnrollment))).scalars().one()
            enrollment.email = "beta-tester@outlook.com"; enrollment.mailbox_key = "beta-tester@outlook.com"
            db.add(alias); await db.commit()
        applied = await self.client.post("/api/beta/applications", headers=self.headers(alias),
            json={"primary_format": "pdf", "motivation": "Rapports trimestriels bilingues"})
        self.assertEqual(applied.status_code, 409)
        self.assertIn("previous_access_active", applied.text)
        async with self.sessions() as db:
            db.add(BetaEnrollment(campaign_id=self.campaign_id, email=alias.email, primary_format="pdf",
                                  limits=DEFAULT_LIMITS))
            await db.commit()
        activated = await self.client.post("/api/beta/activate", headers=self.headers(alias))
        self.assertEqual(activated.status_code, 409)
        self.assertIn("previous_access_active", activated.text)

    async def test_restoring_an_earlier_access_waits_for_the_running_one(self):
        first = await self.activate_tester()
        revoked = await self.client.patch(f"/api/admin/beta/enrollments/{first['id']}", json={"revoke": True},
                                          headers=self.admin_headers)
        self.assertEqual(revoked.status_code, 200, revoked.text)
        await self.open_second_program()
        await self.invite_tester()
        self.assertEqual((await self.client.post("/api/beta/activate", headers=self.tester_headers)).status_code, 200)
        restored = await self.client.patch(f"/api/admin/beta/enrollments/{first['id']}", json={"revoke": False},
                                           headers=self.admin_headers)
        self.assertEqual(restored.status_code, 409)
        self.assertIn("previous_access_active", restored.text)

    async def test_running_access_outranks_a_newer_pending_invitation(self):
        first = await self.activate_tester()
        async with self.sessions() as db:
            db.add(BetaCampaign(id="later", name="Later", state="closed", limits=DEFAULT_LIMITS))
            db.add(BetaEnrollment(campaign_id="later", email=self.tester.email, primary_format="pdf",
                                  limits=DEFAULT_LIMITS, invited_at=beta.now_utc() + timedelta(minutes=1)))
            await db.commit()
        me = (await self.client.get("/api/beta/me", headers=self.tester_headers)).json()
        self.assertEqual((me["enrollment"]["id"], me["enrollment"]["state"]), (first["id"], "active"))

    async def test_invitations_target_the_program_on_screen(self):
        await self.start_campaign()
        draft = await self.client.post("/api/admin/beta/campaigns", json={"name": "Draft", "capacity": 40,
            "first_wave_capacity": 20, "access_days": 30, "limits": DEFAULT_LIMITS}, headers=self.admin_headers)
        into_draft = await self.client.post("/api/admin/beta/invitations", headers=self.admin_headers,
            json={"email": self.tester.email, "campaign_id": draft.json()["id"]})
        self.assertEqual(into_draft.status_code, 409)
        self.assertIn("campaign_closed", into_draft.text)
        missing = await self.client.post("/api/admin/beta/invitations", headers=self.admin_headers,
            json={"email": self.tester.email, "campaign_id": "nope"})
        self.assertEqual(missing.status_code, 404)
        on_screen = await self.client.post("/api/admin/beta/invitations", headers=self.admin_headers,
            json={"email": self.tester.email, "campaign_id": self.campaign_id})
        self.assertEqual(on_screen.status_code, 201, on_screen.text)

    async def test_a_program_past_its_recruitment_window_yields_the_open_slot(self):
        await self.start_campaign()
        first = self.campaign_id
        async with self.sessions() as db:
            (await db.get(BetaCampaign, first)).recruitment_ends_at = beta.now_utc() - timedelta(seconds=1)
            await db.commit()
        draft = await self.client.post("/api/admin/beta/campaigns", json={"name": "Next", "capacity": 40,
            "first_wave_capacity": 20, "access_days": 30, "limits": DEFAULT_LIMITS}, headers=self.admin_headers)
        opened = await self.client.patch(f"/api/admin/beta/campaigns/{draft.json()['id']}", json={"state": "open"},
                                         headers=self.admin_headers)
        self.assertEqual(opened.status_code, 200, opened.text)
        async with self.sessions() as db:
            self.assertEqual((await db.get(BetaCampaign, first)).state, "closed")
        await self.client.patch(f"/api/admin/beta/campaigns/{draft.json()['id']}", json={"state": "closed"},
                                headers=self.admin_headers)
        reopened = await self.client.patch(f"/api/admin/beta/campaigns/{first}", json={"state": "open"},
                                           headers=self.admin_headers)
        self.assertEqual(reopened.status_code, 409)
        self.assertIn("campaign_recruitment_ended", reopened.text)


@unittest.skipUnless(os.getenv("CI") == "true" and
    urlsplit(DATABASE_URL).hostname in ("localhost", "127.0.0.1") and
    urlsplit(DATABASE_URL).path == "/precis", "PostgreSQL concurrency is verified in isolated CI")
class PostgresConcurrencyTests(unittest.IsolatedAsyncioTestCase):
    async def test_simultaneous_imports_cannot_spend_the_last_page_twice(self):
        engine = create_async_engine(DATABASE_URL, poolclass=NullPool)
        sessions = async_sessionmaker(engine, expire_on_commit=False)
        temporary = tempfile.TemporaryDirectory(prefix="precis-beta-race-")
        user_id = enrollment_id = None
        queued = []
        try:
            async with sessions() as db:
                self.assertIsNone(await db.get(BetaCampaign, "pilot"), "Beta concurrency fixture must be isolated")
                campaign = BetaCampaign(id="pilot", name="CI race", state="open", capacity=40,
                    first_wave_capacity=20, access_days=30, limits={**DEFAULT_LIMITS, "pdf": 1},
                    starts_at=beta.now_utc(), recruitment_ends_at=beta.now_utc()+timedelta(days=30))
                user = User(email=f"beta-race-{uuid.uuid4().hex}@example.com", plan="free",
                            email_verified=True, page_credits=7)
                db.add_all([campaign, user]); await db.flush()
                enrollment = BetaEnrollment(campaign_id="pilot", email=user.email, user_id=user.id,
                    primary_format="pdf", limits={**DEFAULT_LIMITS, "pdf": 1},
                    activated_at=beta.now_utc(), expires_at=beta.now_utc()+timedelta(days=30))
                db.add(enrollment); await db.commit()
                user_id, enrollment_id = user.id, enrollment.id
                headers = {"Authorization": "Bearer "+create_access_token(user.id, user.email),
                           "X-API-Key": config.FRONTEND_API_KEY}
            app = create_app()
            async def database():
                async with sessions() as db:
                    yield db
            app.dependency_overrides[get_db] = database
            gate = asyncio.Event()
            calls = 0
            original = beta.check_cost
            async def overlap(*args, **kwargs):
                nonlocal calls
                await original(*args, **kwargs)
                calls += 1
                if calls <= 2:
                    if calls == 2:
                        gate.set()
                    await gate.wait()
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
                with patch.object(beta, "check_cost", side_effect=overlap), \
                     patch.object(translation.scheduler, "submit", side_effect=lambda job_id, *_: queued.append(job_id)), \
                     patch.object(translation, "build_job_paths", side_effect=lambda digest, lang: (
                         str(Path(temporary.name)/digest), str(Path(temporary.name)/digest/lang))):
                    async def request():
                        return await client.post("/api/translate", headers=headers,
                            files={"file": ("race.pdf", pdf())}, data={"use_beta": "1"})
                    responses = await asyncio.wait_for(asyncio.gather(request(), request()), timeout=20)
            self.assertEqual(sorted(r.status_code for r in responses), [200, 402])
            async with sessions() as db:
                self.assertEqual((await beta.totals(db, enrollment_id))["pdf"]["units"], 1)
                self.assertEqual((await db.get(User, user_id)).page_credits, 7)
        finally:
            for job_id in queued:
                jobs.forget(job_id)
            if enrollment_id:
                async with sessions() as db:
                    await db.execute(delete(BetaUsage).where(BetaUsage.enrollment_id == enrollment_id))
                    await db.execute(delete(Document).where(Document.user_id == user_id))
                    await db.execute(delete(BetaEnrollment).where(BetaEnrollment.id == enrollment_id))
                    await db.execute(delete(BetaCampaign).where(BetaCampaign.id == "pilot"))
                    await db.execute(delete(User).where(User.id == user_id))
                    await db.commit()
            await engine.dispose(); temporary.cleanup()

    async def test_simultaneous_openings_leave_a_single_open_program(self):
        engine = create_async_engine(DATABASE_URL, poolclass=NullPool)
        sessions = async_sessionmaker(engine, expire_on_commit=False)
        ids, admin_id = [], None
        try:
            async with sessions() as db:
                self.assertIsNone((await db.execute(select(BetaCampaign).where(BetaCampaign.state == "open"))).first(),
                                  "Beta concurrency fixture must be isolated")
                admin = User(email=f"beta-race-admin-{uuid.uuid4().hex}@example.com", plan="admin", email_verified=True)
                drafts = [BetaCampaign(name=f"CI race {n}", state="draft", capacity=40, first_wave_capacity=20,
                                       access_days=30, limits=DEFAULT_LIMITS) for n in (1, 2)]
                db.add_all([admin, *drafts]); await db.commit()
                ids, admin_id = [row.id for row in drafts], admin.id
                headers = {"Authorization": "Bearer "+create_access_token(admin.id, admin.email),
                           "X-API-Key": config.FRONTEND_API_KEY}
            app = create_app()
            async def database():
                async with sessions() as db:
                    yield db
            app.dependency_overrides[get_db] = database
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
                responses = await asyncio.wait_for(asyncio.gather(*(
                    client.patch(f"/api/admin/beta/campaigns/{campaign_id}", json={"state": "open"}, headers=headers)
                    for campaign_id in ids)), timeout=20)
            self.assertEqual(sorted(r.status_code for r in responses), [200, 409])
            async with sessions() as db:
                opened = (await db.execute(select(BetaCampaign).where(BetaCampaign.state == "open"))).scalars().all()
                self.assertEqual(len(opened), 1)
        finally:
            async with sessions() as db:
                if ids:
                    await db.execute(delete(BetaCampaign).where(BetaCampaign.id.in_(ids)))
                if admin_id:
                    await db.execute(delete(User).where(User.id == admin_id))
                await db.commit()
            await engine.dispose()


if __name__ == "__main__":
    unittest.main()
