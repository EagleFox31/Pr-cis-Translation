"""Un échec en cours de traduction PPTX ne doit pas jeter ce qui est déjà traduit.

LE DÉFAUT QUE CETTE SUITE VERROUILLE
------------------------------------
Les diapositives traduites ne vivaient qu'en mémoire jusqu'à la fin du job :
un échec à la dernière (disque plein, fournisseur coupé) perdait toutes les
autres, déjà payées, et la relance repartait de zéro. Et l'original, écrit
disque plein, restait à 0 octet sans jamais être réécrit : le document
devenait irrécupérable.

On rejoue le VRAI job sur un PPTX synthétique, avec un traducteur factice
(aucun réseau) et la conversion LibreOffice neutralisée (l'aperçu est
best-effort, il n'est pas l'objet de cette suite).

    backend/venv/Scripts/python.exe backend/tests/test_reprise_pptx.py
"""
from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile
import threading
import time

import racine  # noqa: F401  -- met backend/ sur le chemin

from pptx import Presentation  # noqa: E402

from app.services import translation_runner as tr  # noqa: E402
from app.services.jobs import ArretDemande, jobs  # noqa: E402

TEXTES = ["Premier texte", "Deuxieme texte", "Troisieme texte"]


def _pptx(chemin: str) -> bytes:
    prs = Presentation()
    for texte in TEXTES:
        diapo = prs.slides.add_slide(prs.slide_layouts[5])
        diapo.shapes.title.text = texte
    prs.save(chemin)
    with open(chemin, "rb") as f:
        return f.read()


def _cible(texte: str) -> str:
    """Traduction factice, DANS les balises de run : « [[0]]x[[/0]] » → « [[0]]T:x[[/0]] »."""
    return texte.replace("]]", "]]T:", 1)


def _nu(texte: str) -> str:
    return texte.replace("[[0]]", "").replace("[[/0]]", "")


class Traducteur:
    """Préfixe chaque fragment ; échoue sur `echoue_sur` tant qu'il est posé."""

    def __init__(self):
        self.appels: list[str] = []
        self.echoue_sur: str | None = None

    def fiche_document(self, *_a, **_k):
        return ""

    def translate_json(self, chemin, target_lang, progress_callback=None, **_):
        with open(chemin, encoding="utf-8") as f:
            data = json.load(f)
        sources = [_nu(el["text"]) for sd in data["slides"] for el in tr._elements_de(sd)]
        self.appels.extend(sources)
        if self.echoue_sur and any(self.echoue_sur in s for s in sources):
            return False, "fournisseur coupé"
        for sd in data["slides"]:
            for el in tr._elements_de(sd):
                el["translated_text"] = _cible(el["text"])
        sortie = chemin + ".out.json"
        with open(sortie, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False)
        return True, sortie

    def _translate_batch(self, *a, **k):      # passe de cohérence : rien à faire
        return None


def _lancer(dossier: str, octets: bytes, nom: str) -> tuple[str, dict]:
    job_id = jobs.create()
    original = os.path.join(dossier, "original.pptx")
    traduction = os.path.join(dossier, "translated.json")
    sortie = os.path.join(dossier, f"render_{nom}.pptx")
    tr.run_pptx_progressive_job(
        job_id, octets, original, sortie, f"{nom}.pptx",
        os.path.join(dossier, f"partial_{nom}.pdf"), traduction, "en")
    return traduction, jobs.get(job_id) or {}


def _textes_sortie(chemin: str) -> list[str]:
    prs = Presentation(chemin)
    return [sh.text_frame.text for s in prs.slides for sh in s.shapes
            if sh.has_text_frame and sh.text_frame.text]


def main() -> int:
    checks: list[tuple[str, bool, str]] = []

    def ok(nom, cond, detail=""):
        checks.append((nom, bool(cond), detail))

    def _sans_libreoffice(*_a, **_k):
        raise RuntimeError("LibreOffice neutralisé")

    tr.convert_to_pdf_bytes = _sans_libreoffice
    trad = Traducteur()
    tr.ai_translator, tr.ai_active = trad, True

    dossier = tempfile.mkdtemp(prefix="reprise_pptx_")
    try:
        octets = _pptx(os.path.join(dossier, "source.pptx"))
        traduction = os.path.join(dossier, "translated.json")
        reprise = tr._chemin_reprise(traduction)

        # ── 1. Premier essai : la diapo 3 échoue ─────────────────────────
        trad.echoue_sur = "Troisieme"
        _lancer(dossier, octets, "a")
        ok("1. échec : aucune traduction finale", not os.path.exists(traduction))
        gardees = tr._lire_reprise(reprise)
        ok("1. échec : diapos 1 et 2 gardées", sorted(gardees) == [1, 2],
           str(sorted(gardees)))

        # ── 2. Relance : seule la diapo manquante part au modèle ─────────
        trad.echoue_sur, trad.appels = None, []
        evenements: list[tuple[int, str]] = []
        emit = jobs.emit

        def _espion(job_id, type_, data=None, *a, **k):
            if type_ == "page":
                evenements.append((data["page"], data["status"]))
            return emit(job_id, type_, data, *a, **k)

        jobs.emit = _espion
        try:
            _lancer(dossier, octets, "b")
        finally:
            jobs.emit = emit
        # Reprise visible : une diapo entièrement reprise passe « done » sans
        # rejouer ses étapes ; celle qui reste à traduire les rejoue.
        directes = [p for p in (1, 2)
                    if (p, "done") in evenements and (p, "extracting") not in evenements]
        ok("2. relance : diapos reprises affichées traduites d'emblée",
           directes and (3, "extracting") in evenements, str(evenements))
        # Rien de ce qui était déjà traduit ne repart au modèle, et la diapo
        # en échec, elle, y part bien.
        ok("2. relance : seuls les fragments manquants retraduits",
           "Troisieme texte" in trad.appels
           and not {"Premier texte", "Deuxieme texte"} & set(trad.appels)
           and len(trad.appels) == len(set(trad.appels)), str(trad.appels))
        # Le gabarit (texte « Click to edit… ») est pris par une diapo ou une
        # autre selon la course : il doit être repris quel que soit son porteur.
        ok("2. relance : traduction finale écrite", os.path.isfile(traduction))
        if os.path.isfile(traduction):
            with open(traduction, encoding="utf-8") as f:
                final = json.load(f)
            paires = [(_nu(el["text"]), el.get("translated_text"))
                      for sd in final["slides"] for el in tr._elements_de(sd)]
            ok("2. relance : tout traduit, source conservée",
               paires and all(c == _cible(f"[[0]]{s}[[/0]]") for s, c in paires)
               and {s for s, _ in paires} >= set(TEXTES), str(paires))
        ok("2. relance : fichier de reprise retiré", not os.path.exists(reprise))
        sortie = os.path.join(dossier, "render_b.pptx")
        ok("2. relance : PPTX final traduit",
           os.path.isfile(sortie) and _textes_sortie(sortie) == ["T:" + t for t in TEXTES],
           str(_textes_sortie(sortie)) if os.path.isfile(sortie) else "absent")

        # ── 3. Reprise périmée : un texte source qui ne correspond plus ──
        os.remove(traduction)
        perimee = {1: {"text_elements": [{"id": "x", "text": "autre chose",
                                          "translated_text": "FAUX"}]}}
        tr._ecrire_reprise(reprise, perimee)
        trad.appels = []
        _lancer(dossier, octets, "c")
        ok("3. reprise périmée : rien n'est repris à tort",
           set(TEXTES) <= set(trad.appels), str(trad.appels))

        # ── 4. Original à 0 octet : réécrit, pas réutilisé ───────────────
        original = os.path.join(dossier, "original.pptx")
        open(original, "wb").close()
        tr._sauver_original(original, octets)
        ok("4. original vide réécrit", os.path.getsize(original) == len(octets))
        tr._sauver_original(original, b"")
        ok("4. original complet gardé sans octets neufs",
           os.path.getsize(original) == len(octets))
        open(original, "wb").close()
        try:
            tr._sauver_original(original, b"")
            ok("4. original vide sans octets neufs refusé", False)
        except ValueError:
            ok("4. original vide sans octets neufs refusé", True)

        # ── 5. Pause demandée : aucune diapo ne part au modèle ───────────
        if os.path.exists(traduction):
            os.remove(traduction)
        trad.appels = []
        job_id = jobs.create()
        jobs.demander_arret(job_id, "paused")
        try:
            tr.run_pptx_progressive_job(
                job_id, octets, original, os.path.join(dossier, "render_d.pptx"),
                "d.pptx", os.path.join(dossier, "partial_d.pdf"), traduction, "en")
            ok("5. pause : l'arrêt remonte au planificateur", False)
        except ArretDemande as a:
            ok("5. pause : l'arrêt remonte au planificateur", a.mode == "paused", a.mode)
        ok("5. pause : aucun fragment traduit", not trad.appels, str(trad.appels))
        ok("5. pause : pas d'état « erreur »",
           (jobs.get(job_id) or {}).get("state") != "error",
           str((jobs.get(job_id) or {}).get("state")))
        jobs.forget(job_id)

        # ── 5b. Arrêt pendant la traduction : diapos en vol abandonnées ──
        # Le modèle « ne répond pas » (4 s) : l'arrêt n'attend pas son retour,
        # et la réponse tardive n'est ni gardée ni mise en reprise.
        for f in (traduction, reprise):
            if os.path.exists(f):
                os.remove(f)
        entree = threading.Event()
        vrai = trad.translate_json

        def _lent(*a, **k):
            entree.set()
            time.sleep(4)
            return vrai(*a, **k)
        trad.translate_json = _lent
        job_id = jobs.create()
        issue: dict = {}

        def _job():
            try:
                tr.run_pptx_progressive_job(
                    job_id, octets, original, os.path.join(dossier, "render_g.pptx"),
                    "g.pptx", os.path.join(dossier, "partial_g.pdf"), traduction, "en")
            except ArretDemande as a:
                issue["mode"] = a.mode
            issue["fin"] = time.monotonic()
        fil = threading.Thread(target=_job)
        fil.start()
        entree.wait(30)
        debut = time.monotonic()
        jobs.demander_arret(job_id, "cancelled")
        fil.join(30)
        ok("5b. arrêt en vol : remonte sans attendre le modèle",
           issue.get("mode") == "cancelled" and issue.get("fin", 99) - debut < 2.5,
           str(issue))
        time.sleep(4.5)                   # la réponse tardive arrive
        ok("5b. arrêt en vol : diapo tardive ni gardée ni en reprise",
           not tr._lire_reprise(reprise) and not os.path.exists(traduction),
           str(tr._lire_reprise(reprise)))
        trad.translate_json = vrai
        jobs.forget(job_id)

        # ── 6. Document arrêté : le fichier de reprise se livre tel quel ──
        # Les diapositives faites traduites, les autres en langue d'origine.
        for f in (traduction, reprise):
            if os.path.exists(f):
                os.remove(f)
        trad.echoue_sur = "Deuxieme"
        _lancer(dossier, octets, "e")
        trad.echoue_sur = None
        rendu = os.path.join(dossier, "partiel.pptx")
        with open(rendu, "wb") as f:
            f.write(tr.render_translation_bytes(original, reprise, "pptx", "en"))
        textes = _textes_sortie(rendu)
        ok("6. partiel : la diapositive faite est livrée traduite",
           "T:Premier texte" in textes, str(textes))
        ok("6. partiel : la diapositive manquée reste d'origine",
           "Deuxieme texte" in textes, str(textes))

        # ── 7. Reprise remplacée par la traduction finale d'un autre compte ──
        # Le document arrêté n'en reçoit que ses diapositives payées.
        _lancer(dossier, octets, "f")
        ok("7. traduction finale écrite, reprise retirée",
           os.path.isfile(traduction) and not os.path.exists(reprise))
        with open(rendu, "wb") as f:
            f.write(tr.render_translation_bytes(original, traduction, "pptx", "en",
                                                only_pages={1}))
        textes = _textes_sortie(rendu)
        ok("7. finale bornée : la diapositive payée est traduite",
           "T:Premier texte" in textes, str(textes))
        ok("7. finale bornée : les autres restent d'origine",
           "Deuxieme texte" in textes and "T:Deuxieme texte" not in textes, str(textes))
    finally:
        shutil.rmtree(dossier, ignore_errors=True)

    for nom, passe, detail in checks:
        print(("OK   " if passe else "ECHEC"), nom, ("" if passe else detail))
    n_ok = sum(1 for _, p, _ in checks if p)
    print(f"{n_ok}/{len(checks)}")
    return 0 if n_ok == len(checks) else 1


if __name__ == "__main__":
    sys.exit(main())
