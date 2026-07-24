"""Prévisualisation PPTX — les trois défauts qui la rendaient lente et fausse.

Ce que ce fichier verrouille, et POURQUOI chaque test existe :

1. `build_partial_pptx` — la sélection des slides à garder excluait par PRÉFIXE
   de chemin (`startswith("ppt/slides/slide1")`), ce qui emportait aussi slide12
   et slide19. Le défaut était invisible tant qu'on ne coupait qu'une queue
   contiguë ; `only_slides` (garder une slide isolée) le rend immédiat.
   ⚠ L'aperçu progressif NE s'en sert PAS : convertir slide par slide a été
   mesuré 17× plus lent que convertir le deck entier (le coût de LibreOffice est
   son DÉMARRAGE, ~7 s, pas le rendu). Cf. `_write_pptx_pdf_partial`.

2. ISOLATION du moteur — `app.pptx_engine` était UNE instance pour tout le
   processus, avec UN dossier temporaire. Deux réinjections simultanées y
   écrivaient donc dans les mêmes fichiers XML : on demandait l'anglais, on
   recevait le portugais. Le test lance les deux EN MÊME TEMPS et vérifie
   qu'aucune sortie ne contient un mot de l'autre.

3. CACHE de l'aperçu non-PDF — le rendu d'un PPTX est intégral par nature
   (réinjection du deck + LibreOffice) ; il était pourtant jeté dès que la
   requête portait un `?page=N`, que le frontend envoie TOUJOURS. Le test
   COMPTE les rendus : deux pages consultées ne doivent en coûter qu'un.

Le PPTX de test est SYNTHÉTIQUE (python-pptx), jamais un document qui aurait
révélé le bug : un correctif ne vaut que s'il se prouve sur un document
quelconque.

Exécution :  backend/venv/Scripts/python.exe backend/test_pptx_preview.py
Les tests 1 et 2 sont autonomes ; le 3 requiert PostgreSQL.
"""
from __future__ import annotations
import asyncio
import json
import os
import shutil
import sys
import tempfile
import threading
import uuid
import zipfile

import racine  # noqa: F401  -- met backend/ sur le chemin

from lxml import etree                                    # noqa: E402
from pptx import Presentation                             # noqa: E402
from pptx.util import Inches                              # noqa: E402

from engines.pptx.engine import PPTXTranslatorEngine, NAMESPACES  # noqa: E402

NB_SLIDES = 12


def _deck_synthetique(path: str, n: int = NB_SLIDES) -> None:
    """Un deck de `n` slides, chacune portant un texte qui l'identifie.

    12 slides et non 3 : c'est à partir de 10 que « slide1 » devient un préfixe
    de « slide12 ». Un deck court laisserait passer le défaut n°1.
    """
    prs = Presentation()
    vide = prs.slide_layouts[6]              # disposition sans placeholder
    for i in range(1, n + 1):
        slide = prs.slides.add_slide(vide)
        zone = slide.shapes.add_textbox(Inches(1), Inches(1),
                                        Inches(6), Inches(1))
        zone.text_frame.text = f"CONTENU DE LA SLIDE {i}"
    prs.save(path)


def _slides_du_zip(path: str) -> set[int]:
    """Numéros des slides RÉELLEMENT présentes dans le paquet."""
    import re
    with zipfile.ZipFile(path) as zf:
        noms = zf.namelist()
    return {int(m.group(1)) for m in
            (re.fullmatch(r"ppt/slides/slide(\d+)\.xml", nom) for nom in noms)
            if m}


def _slides_declarees(path: str) -> int:
    """Nombre de slides déclarées dans presentation.xml — ce que LibreOffice
    rendra réellement. Un fichier présent mais non déclaré ne produit pas de
    page ; l'inverse produit un paquet cassé."""
    with zipfile.ZipFile(path) as zf:
        tree = etree.fromstring(zf.read("ppt/presentation.xml"))
    lst = tree.xpath("//p:sldIdLst", namespaces=NAMESPACES)
    return len(lst[0]) if lst else 0


def _textes(path: str) -> str:
    """Tout le texte des slides du paquet, concaténé."""
    out = []
    with zipfile.ZipFile(path) as zf:
        for nom in zf.namelist():
            if nom.startswith("ppt/slides/slide") and nom.endswith(".xml"):
                root = etree.fromstring(zf.read(nom))
                out += [t.text or "" for t in
                        root.xpath(".//a:t", namespaces=NAMESPACES)]
    return " ".join(out)


# ── 1. Sélection de slides ───────────────────────────────────────────────────

def test_selection(ok) -> None:
    with tempfile.TemporaryDirectory() as td:
        src = os.path.join(td, "deck.pptx")
        _deck_synthetique(src)

        eng = PPTXTranslatorEngine()
        eng._extract_zip(src)
        try:
            # a) LA slide isolée de plus haut numéro. Avec l'exclusion par
            #    préfixe, « slide1 » emportait « slide12 » : le paquet sortait
            #    VIDE de toute slide.
            out = os.path.join(td, "seule12.pptx")
            eng.build_partial_pptx(out, NB_SLIDES, only_slides={12})
            ok("SELECTION  only_slides={12} garde bien la slide 12",
               _slides_du_zip(out) == {12}, f"présentes={_slides_du_zip(out)}")
            ok("SELECTION  only_slides={12} déclare exactement 1 slide",
               _slides_declarees(out) == 1, f"déclarées={_slides_declarees(out)}")
            ok("SELECTION  only_slides={12} ne contient que SON texte",
               "SLIDE 12" in _textes(out) and "SLIDE 1 " not in _textes(out))

            # b) Une slide du milieu.
            out = os.path.join(td, "seule3.pptx")
            eng.build_partial_pptx(out, NB_SLIDES, only_slides={3})
            ok("SELECTION  only_slides={3} garde bien la slide 3",
               _slides_du_zip(out) == {3}, f"présentes={_slides_du_zip(out)}")

            # c) Compatibilité : sans `only_slides`, c'est toujours 1..up_to.
            out = os.path.join(td, "jusqua5.pptx")
            eng.build_partial_pptx(out, 5)
            ok("SELECTION  up_to_slide=5 garde 1..5",
               _slides_du_zip(out) == {1, 2, 3, 4, 5},
               f"présentes={_slides_du_zip(out)}")
            ok("SELECTION  up_to_slide=5 déclare 5 slides",
               _slides_declarees(out) == 5, f"déclarées={_slides_declarees(out)}")

            # d) Le deck complet reste intact (c'est le fichier LIVRÉ).
            out = os.path.join(td, "tout.pptx")
            eng.build_partial_pptx(out, NB_SLIDES)
            ok("SELECTION  deck complet : les 12 slides sont là",
               _slides_du_zip(out) == set(range(1, NB_SLIDES + 1)),
               f"présentes={_slides_du_zip(out)}")
        finally:
            eng._cleanup_temp()


# ── 2. Isolation des moteurs ─────────────────────────────────────────────────

def _json_traduit(extraction: dict, marqueur: str) -> dict:
    """Recopie l'extraction en remplaçant chaque texte par `marqueur`,
    balises de runs comprises (le moteur réinjecte run par run)."""
    import re
    data = json.loads(json.dumps(extraction))
    for slide in data.get("slides", []):
        for el in slide.get("text_elements", []):
            el["translated_text"] = re.sub(
                r"\[\[(\d+)\]\].*?\[\[/\1\]\]",
                lambda m: f"[[{m.group(1)}]]{marqueur}[[/{m.group(1)}]]",
                el["text"], flags=re.DOTALL)
    return data


def test_isolation(ok) -> None:
    """Deux réinjections SIMULTANÉES ne doivent pas se mélanger.

    C'est le bug « je clique sur la traduction anglaise, je vois la
    portugaise » : les deux aperçus partageaient le dossier temporaire du
    moteur unique du processus.
    """
    with tempfile.TemporaryDirectory() as td:
        src = os.path.join(td, "deck.pptx")
        _deck_synthetique(src)

        sonde = PPTXTranslatorEngine()
        extraction, _ = sonde.extract_text(
            src, os.path.join(td, "extraction.json"))
        sonde._cleanup_temp()
        if not extraction:
            ok("ISOLATION  extraction du deck synthétique", False,
               "aucun texte extrait")
            return

        langues = {"ANGLAISXX": os.path.join(td, "en.pptx"),
                   "PORTUGAIS": os.path.join(td, "pt.pptx")}
        for marqueur in langues:
            with open(os.path.join(td, f"{marqueur}.json"), "w",
                      encoding="utf-8") as f:
                json.dump(_json_traduit(extraction, marqueur), f,
                          ensure_ascii=False)

        # Départ SIMULTANÉ : la barrière garantit que les deux réinjections se
        # recouvrent réellement. Sans elle, un test séquentiel passerait même
        # avec le moteur partagé.
        barriere = threading.Barrier(len(langues))
        erreurs: dict[str, str] = {}

        def _injecter(marqueur: str, sortie: str) -> None:
            eng = PPTXTranslatorEngine()          # ← UN MOTEUR PAR OPÉRATION
            barriere.wait()
            reussi, msg = eng.inject_translation(
                src, os.path.join(td, f"{marqueur}.json"), sortie)
            if not reussi:
                erreurs[marqueur] = msg

        fils = [threading.Thread(target=_injecter, args=(m, s))
                for m, s in langues.items()]
        for t in fils:
            t.start()
        for t in fils:
            t.join(timeout=120)

        ok("ISOLATION  les deux réinjections aboutissent",
           not erreurs, str(erreurs))

        for marqueur, sortie in langues.items():
            autres = [m for m in langues if m != marqueur]
            if not os.path.isfile(sortie):
                ok(f"ISOLATION  {marqueur} : fichier produit", False, "absent")
                continue
            texte = _textes(sortie)
            ok(f"ISOLATION  {marqueur} : contient bien SA traduction",
               marqueur in texte, f"extrait={texte[:80]!r}")
            for autre in autres:
                ok(f"ISOLATION  {marqueur} : aucune trace de {autre}",
                   autre not in texte, f"extrait={texte[:120]!r}")


# ── 1-bis. Parties partagées : layouts et masques ────────────────────────────

def _deck_meme_layout(path: str, n: int = 5):
    """`n` diapositives qui PARTAGENT le même slideLayout.

    Le modèle par défaut de python-pptx porte du texte dans ses layouts
    (« Click to edit Master title style » …) : c'est exactement la situation
    d'un masque d'entreprise, sans avoir à en fabriquer un.
    """
    prs = Presentation()
    layout = prs.slide_layouts[1]
    for _ in range(n):
        prs.slides.add_slide(layout)
    prs.save(path)


def test_layouts_partages(ok) -> None:
    """Un layout partagé n'est extrait, traduit et injecté QU'UNE FOIS.

    Il était extrait une fois PAR DIAPOSITIVE qui le référence : autant de
    traductions payées pour le même texte, injectées tour à tour dans le même
    fichier — la dernière écrasant les autres. Et en progressif, la deuxième
    extraction relisait un layout DÉJÀ traduit : on traduisait une traduction.
    """
    with tempfile.TemporaryDirectory() as td:
        src = os.path.join(td, "partage.pptx")
        _deck_meme_layout(src, n=5)

        eng = PPTXTranslatorEngine()
        eng._extract_zip(src)
        try:
            porteuses, ids = [], []
            for sn in range(1, eng.slide_count() + 1):
                sd, _ = eng.extract_slide(sn)
                if not sd:
                    continue
                for bloc in sd.get("layout_elements", []):
                    porteuses.append(sn)
                    ids += [e["id"] for e in bloc["text_elements"]]
                # Le job injecte APRÈS chaque extraction : on le reproduit, car
                # c'est ce qui rendait la 2e extraction fautive.
                eng.inject_slide(sn, {e["id"]: e["text"]
                                      for e in _tous_les_elements(sd)})

            ok("LAYOUT  une seule diapositive porte le layout partagé",
               len(porteuses) == 1, f"porteuses={porteuses}")
            ok("LAYOUT  du texte de layout a bien été extrait", bool(ids))
            ok("LAYOUT  son identité ne dépend pas de la diapositive",
               all(i.startswith("slide0_layout_") for i in ids), str(ids[:3]))
        finally:
            eng._cleanup_temp()

        # L'injection atteint-elle vraiment le layout ?
        eng = PPTXTranslatorEngine()
        eng._extract_zip(src)
        try:
            carte = {}
            for sn in range(1, eng.slide_count() + 1):
                sd, _ = eng.extract_slide(sn)
                if not sd:
                    continue
                for el in _tous_les_elements(sd):
                    carte[el["id"]] = el["text"].replace(
                        "Master", "ZZTRADUITZZ")
                eng.inject_slide(sn, carte)
            sortie = os.path.join(td, "injecte.pptx")
            eng.build_partial_pptx(sortie, eng.slide_count())
        finally:
            eng._cleanup_temp()

        with zipfile.ZipFile(sortie) as zf:
            layouts = [n for n in zf.namelist()
                       if n.startswith("ppt/slideLayouts/slideLayout")]
            touche = [n for n in layouts if b"ZZTRADUITZZ" in zf.read(n)]
        ok("LAYOUT  la traduction atteint bien le fichier de layout",
           len(touche) >= 1, f"layouts modifiés={touche}")

    # Compatibilité : les traductions produites AVANT ce changement nomment les
    # paragraphes `slide{N}_layout_…`. Sans repli, le layout resterait en
    # langue source sur tous les documents déjà traduits.
    with tempfile.TemporaryDirectory() as td:
        src = os.path.join(td, "compat.pptx")
        _deck_meme_layout(src, n=2)
        eng = PPTXTranslatorEngine()
        eng._extract_zip(src)
        try:
            slides_dir = eng._get_temp_dir() / "ppt" / "slideLayouts"
            cible = sorted(slides_dir.glob("slideLayout*.xml"))[0]
            ancienne = {f"slide3_layout_{cible.stem}_p0":
                        "[[0]]ANCIENNE IDENTITE[[/0]]"}
            touche = eng._inject_in_xml(cible, PPTXTranslatorEngine.PART_PARTAGEE,
                                        f"layout_{cible.stem}", ancienne,
                                        num_alt=(3,))
            ok("LAYOUT  une ancienne identité est encore retrouvée",
               touche and b"ANCIENNE IDENTITE" in cible.read_bytes())
        finally:
            eng._cleanup_temp()


def _tous_les_elements(slide_data):
    yield from slide_data.get("text_elements", [])
    for groupe in ("diagram_elements", "chart_elements", "layout_elements"):
        for bloc in slide_data.get(groupe, []):
            yield from bloc.get("text_elements", [])
    yield from slide_data.get("excel_elements", [])


# ── 2-bis. Cohérence terminologique du document ──────────────────────────────

def test_coherence(ok) -> None:
    """La passe qui reprend un terme traduit ailleurs et oublié ici.

    Le modèle traduit slide par slide, sans mémoire d'une slide à l'autre : il
    rend « Gerbeur » par « Stacker » sur deux diapositives et le laisse en
    français sur la troisième. Aucune règle de FORME ne l'attrape — le fragment
    est parfaitement balisé. Seul le DOCUMENT prouve que le terme est
    traduisible.

    Ce qui compte ici n'est pas seulement que la reprise ait lieu, mais qu'elle
    soit CHIRURGICALE : un seul fragment redemandé, une seule slide réinjectée,
    et les traductions déjà correctes laissées intactes.
    """
    from app.services import translation_runner as appmod

    appels = []

    class _FauxTraducteur:
        def _translate_batch(self, batch, target_lang, cb=None, **kw):
            appels.append([b["id"] for b in batch])
            for b in batch:
                if "consigne" not in b:
                    raise AssertionError("consigne de cohérence absente")
                b["translated_text"] = b["translated_text"].replace(
                    "Gerbeur", "Stacker")
            return True

    class _FauxMoteur:
        def __init__(self):
            self.reinjections = []

        def inject_slide(self, sn, tmap):
            self.reinjections.append(sn)

    slides = {
        1: {"text_elements": [{"id": "a", "text": "[[0]]Découverte du Gerbeur[[/0]]",
                               "translated_text": "[[0]]Discovering the Stacker[[/0]]"}]},
        2: {"text_elements": [{"id": "b", "text": "[[0]]Manutention du Gerbeur[[/0]]",
                               "translated_text": "[[0]]Moving the Stacker[[/0]]"}]},
        3: {"text_elements": [{"id": "c", "text": "[[0]]fonctionnement du Gerbeur[[/0]]",
                               "translated_text": "[[0]]operation of the Gerbeur[[/0]]"}]},
    }
    vrai = appmod.ai_translator
    appmod.ai_translator = _FauxTraducteur()
    moteur = _FauxMoteur()
    try:
        n = appmod._coherence_document(slides, moteur, "en", None)
    finally:
        appmod.ai_translator = vrai

    ok("COHERENCE  un seul fragment est repris", n == 1, f"n={n}")
    ok("COHERENCE  seul le fragment fautif est redemandé",
       appels == [["c"]], str(appels))
    ok("COHERENCE  seule la slide touchée est réinjectée",
       moteur.reinjections == [3], str(moteur.reinjections))
    ok("COHERENCE  le terme oublié est corrigé",
       "Stacker" in slides[3]["text_elements"][0]["translated_text"]
       and "Gerbeur" not in slides[3]["text_elements"][0]["translated_text"],
       slides[3]["text_elements"][0]["translated_text"])
    ok("COHERENCE  les traductions déjà correctes sont intactes",
       slides[1]["text_elements"][0]["translated_text"]
       == "[[0]]Discovering the Stacker[[/0]]")
    ok("COHERENCE  la consigne de reprise ne reste pas dans le JSON",
       "consigne" not in slides[3]["text_elements"][0])

    # Un document COHÉRENT ne doit coûter aucun appel.
    appels.clear()
    appmod.ai_translator = _FauxTraducteur()
    try:
        n = appmod._coherence_document(
            {1: slides[1], 2: slides[2]}, _FauxMoteur(), "en", None)
    finally:
        appmod.ai_translator = vrai
    ok("COHERENCE  un document cohérent ne coûte aucun appel",
       n == 0 and appels == [], f"n={n} appels={appels}")


# ── 3. Cache de l'aperçu non-PDF ─────────────────────────────────────────────

async def test_cache_apercu(ok) -> None:
    """Consulter deux pages d'un aperçu PPTX ne doit coûter QU'UN rendu.

    On COMPTE les rendus au lieu de regarder le fichier de cache : c'est la
    dépense qu'on veut supprimer, et un test qui la compte ne peut pas être
    satisfait par un cache écrit puis ignoré.
    """
    from sqlalchemy.ext.asyncio import (create_async_engine,
                                        async_sessionmaker, AsyncSession)
    from sqlalchemy.pool import NullPool
    from app.core.database import DATABASE_URL
    from app.models import User, Document
    from app.services import render_cache
    # On substitue dans le module qui APPELLE (la route), pas dans celui qui
    # definit : `documents` importe ces deux noms et en garde sa propre
    # reference. Remplacer la definition d'origine ne changerait rien a ce que
    # la route execute -- le test compterait zero appel sans rien prouver.
    from app.api import documents as appmod
    from app.api.documents import STORAGE_BASE, preview_document

    moteur = create_async_engine(DATABASE_URL, poolclass=NullPool)
    session = async_sessionmaker(moteur, class_=AsyncSession,
                                 expire_on_commit=False)

    racine = os.path.join(STORAGE_BASE, f"_test_prev_{uuid.uuid4().hex[:8]}")
    lang_dir = os.path.join(racine, "en")
    os.makedirs(lang_dir, exist_ok=True)
    orig = os.path.join(racine, "original.pptx")
    trad = os.path.join(lang_dir, "translated.json")
    _deck_synthetique(orig, n=3)
    with open(trad, "w", encoding="utf-8") as f:
        f.write('{"slides": []}')

    # On substitue le rendu et la conversion : ce test mesure la POLITIQUE de
    # cache, pas LibreOffice. Chaque appel s'annonce.
    n_rendus = [0]
    n_conversions = [0]
    vrai_rendu = appmod.render_translation_bytes
    vraie_conv = appmod.convert_to_pdf_bytes

    def _rendu(*a, **k):
        n_rendus[0] += 1
        with open(orig, "rb") as f:
            return f.read()

    def _conv(*a, **k):
        n_conversions[0] += 1
        return b"%PDF-1.7\n% apercu\n"

    appmod.render_translation_bytes = _rendu
    appmod.convert_to_pdf_bytes = _conv

    u = User(email=f"prev-{uuid.uuid4().hex[:10]}@example.com",
             plan="starter", email_verified=True)
    try:
        async with session() as db:
            db.add(u)
            await db.commit()
            d = Document(user_id=u.id, original_name="presentation.pptx",
                         source_lang="auto", target_lang="en",
                         original_path=orig, translated_path=trad,
                         size_bytes=1000, status="done", paid=True)
            db.add(d)
            await db.commit()
            doc_id, user_id = d.id, u.id

        async with session() as db:
            user = await db.get(User, user_id)

            r1 = await preview_document(doc_id, user=user, db=db, page=1)
            ok("CACHE  page 1 : rend un PDF",
               bytes(r1.body[:5]).startswith(b"%PDF"))
            ok("CACHE  page 1 : annonce le document COMPLET (X-Render: full)",
               r1.headers.get("X-Render") == "full",
               f"X-Render={r1.headers.get('X-Render')}")
            ok("CACHE  page 1 : un seul rendu",
               n_rendus[0] == 1, f"rendus={n_rendus[0]}")

            r2 = await preview_document(doc_id, user=user, db=db, page=2)
            ok("CACHE  page 2 : AUCUN rendu supplémentaire",
               n_rendus[0] == 1, f"rendus={n_rendus[0]}")
            ok("CACHE  page 2 : aucune conversion LibreOffice supplémentaire",
               n_conversions[0] == 1, f"conversions={n_conversions[0]}")
            ok("CACHE  page 2 : rend le même document",
               bytes(r2.body) == bytes(r1.body))
            ok("CACHE  le rendu est bien conservé sur disque",
               render_cache.cache_valid(trad) is not None)

            # ── Le panneau SOURCE ────────────────────────────────────────
            # Il restait blanc : le client téléchargeait le PPTX natif puis le
            # RENVOYAIT au serveur pour conversion. `?as=pdf` sert directement
            # ce que le serveur a déjà. Sans le paramètre, on doit toujours
            # rendre le NATIF — la relance d'une traduction en dépend.
            from app.api.documents import original_document

            natif = await original_document(doc_id, user=user, db=db)
            ok("SOURCE  sans `as` : rend le format NATIF (PK…)",
               bytes(natif.body[:2]) == b"PK",
               f"head={bytes(natif.body[:4])!r}")

            en_pdf = await original_document(doc_id, user=user, db=db,
                                             as_="pdf")
            ok("SOURCE  `?as=pdf` : rend un PDF",
               bytes(en_pdf.body[:5]).startswith(b"%PDF"),
               f"head={bytes(en_pdf.body[:5])!r}")
            ok("SOURCE  `?as=pdf` : annonce bien application/pdf",
               en_pdf.media_type == "application/pdf",
               f"media_type={en_pdf.media_type}")
    finally:
        appmod.render_translation_bytes = vrai_rendu
        appmod.convert_to_pdf_bytes = vraie_conv
        async with session() as db:
            obj = await db.get(User, user_id if 'user_id' in dir() else u.id)
            if obj:
                await db.delete(obj)
            await db.commit()
        await moteur.dispose()
        shutil.rmtree(racine, ignore_errors=True)


# ── Exécution ────────────────────────────────────────────────────────────────

def main() -> int:
    checks: list[tuple[str, bool, str]] = []

    def ok(nom, cond, detail=""):
        checks.append((nom, bool(cond), detail))

    test_selection(ok)
    test_layouts_partages(ok)
    test_isolation(ok)
    test_coherence(ok)

    if "--sans-db" not in sys.argv:
        try:
            asyncio.run(test_cache_apercu(ok))
        except Exception as e:                      # base indisponible
            print(f"\n(test 3 ignoré : {type(e).__name__}: {e})")

    print()
    n_ok = sum(1 for _, c, _ in checks if c)
    for nom, cond, detail in checks:
        etat = "  OK  " if cond else " FAIL "
        ligne = f"{etat}  {nom}"
        if not cond and detail:
            ligne += f"   [{detail}]"
        print(ligne)
    print(f"\n{n_ok}/{len(checks)}")
    return 0 if n_ok == len(checks) else 1


if __name__ == "__main__":
    sys.exit(main())
